"""Bring a firm's products in from a spreadsheet, with a template to fill in.

Backlog 46. A firm moving from Tally, Excel or another ERP brings its item
master as a file. The rules the import keeps:

* **Every row is checked before anything is written**, and every problem is
  reported with its row number and column -- not only the first, because a
  3,000-row file that fails one row at a time is re-run 3,000 times.
* **All or nothing.** Each row goes through ``stage_product`` or
  ``stage_update_product`` -- the same guards and audit writes as the form --
  and the file commits once, or not at all.
* **References are matched, never guessed.** A category, unit or tax group is
  found by code or by name, ignoring case; one that is not there is reported.
* **Headings are matched ignoring case, spaces and punctuation**, with the
  names other software commonly exports ("Item Code", "HSN/SAC", "Sale
  Price"), so an export needs little editing.
* **Update by code is an option.** With ``existing="update"`` a row whose
  code is already a product updates it, and a blank cell leaves that field
  alone -- so a migration can be corrected and re-run.
"""

# ruff: noqa: D102, D107

import csv
import io
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import TYPE_CHECKING, Literal
from uuid import UUID

from pydantic import ValidationError as PydanticValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.exceptions import ApplicationError, ValidationError
from app.products.models import Product, ProductCategory
from app.products.schemas import ProductCreate, ProductUpdate
from app.products.schemas.product import ProductStatus, ProductType
from app.tax.models import TaxProfile
from app.uom.models import Uom

if TYPE_CHECKING:
    from app.products.services.product_service import ProductService

ExistingRows = Literal["refuse", "update"]


@dataclass(frozen=True)
class Column:
    """One template column: its heading, the names it also answers to, and help."""

    heading: str
    aliases: tuple[str, ...]
    required: bool
    takes: str
    example: str


#: The template, in the order it is laid out. The headings are what the notes
#: sheet explains and what every error names.
COLUMNS: tuple[Column, ...] = (
    Column(
        "Code",
        ("productcode", "itemcode", "sku", "partno", "partnumber"),
        True,
        "Unique product code: letters, digits, - and _ (made upper case).",
        "RICE-5KG",
    ),
    Column(
        "Name",
        ("productname", "itemname"),
        True,
        "Product name as it prints on documents.",
        "Basmati Rice 5 kg",
    ),
    Column(
        "Type",
        ("producttype", "itemtype"),
        False,
        "One of: "
        + ", ".join(item.value for item in ProductType)
        + ". Blank means STOCK_ITEM.",
        "STOCK_ITEM",
    ),
    Column(
        "Category",
        ("categorycode", "categoryname", "stockgroup", "group", "itemgroup"),
        False,
        "An existing category, by code or name (see the Lists sheet).",
        "GROCERY",
    ),
    Column(
        "SubCategory",
        ("subcategorycode", "subcategoryname", "subgroup"),
        False,
        "An existing sub category of that category, by code or name.",
        "",
    ),
    Column(
        "Unit",
        ("uom", "baseunit", "baseuom", "units", "unitofmeasure"),
        False,
        "An existing unit, by code, name or symbol (see the Lists sheet).",
        "PCS",
    ),
    Column(
        "HSN",
        ("hsnsac", "hsncode", "saccode", "sac"),
        False,
        "HSN or SAC code.",
        "10063020",
    ),
    Column(
        "TaxGroup",
        ("taxprofile", "taxprofilegroupcode", "gstgroup", "taxcategory"),
        False,
        "An active tax profile group code (see the Lists sheet).",
        "GST5",
    ),
    Column("Brand", ("make", "manufacturer"), False, "Brand name.", "Acme"),
    Column("Model", ("modelno", "modelnumber"), False, "Model.", ""),
    Column(
        "Barcode",
        ("ean", "upc", "gtin"),
        False,
        "Barcode; must be unique among the firm's products.",
        "",
    ),
    Column("ShortName", ("alias",), False, "Short name, 100 characters.", ""),
    Column("Description", ("details",), False, "Longer description.", ""),
    Column(
        "PurchasePrice",
        ("costprice", "cost", "purchaserate", "buyingprice"),
        False,
        "Number, no currency sign; commas are ignored.",
        "400",
    ),
    Column(
        "SellingPrice",
        ("saleprice", "salesprice", "salerate", "sellingrate", "rate", "price"),
        False,
        "Number; must not exceed MRP.",
        "450",
    ),
    Column("MRP", ("maximumretailprice",), False, "Number.", "499"),
    Column(
        "Status",
        (),
        False,
        "One of: "
        + ", ".join(item.value for item in ProductStatus)
        + ". Blank means ACTIVE.",
        "ACTIVE",
    ),
    Column(
        "TrackBatch",
        ("batchwise", "maintainbatches"),
        False,
        "Yes or No.",
        "No",
    ),
    Column("TrackExpiry", ("expirytracking",), False, "Yes or No.", "No"),
    Column(
        "TrackSerial",
        ("serialnumbers",),
        False,
        "Yes or No.",
        "No",
    ),
    Column(
        "AllowNegativeStock",
        ("negativestock",),
        False,
        "Yes or No.",
        "No",
    ),
    Column("Remarks", ("notes", "note"), False, "Free text.", ""),
)

#: The write-schema field each heading fills, for naming a schema refusal.
_FIELD_HEADINGS: dict[str, str] = {
    "code": "Code",
    "name": "Name",
    "product_type": "Type",
    "category_id": "Category",
    "sub_category_id": "SubCategory",
    "base_uom_id": "Unit",
    "unit": "Unit",
    "hsn_sac": "HSN",
    "tax_profile_group_code": "TaxGroup",
    "brand": "Brand",
    "model": "Model",
    "barcode": "Barcode",
    "short_name": "ShortName",
    "description": "Description",
    "purchase_price": "PurchasePrice",
    "selling_price": "SellingPrice",
    "mrp": "MRP",
    "status": "Status",
    "track_batch": "TrackBatch",
    "track_expiry": "TrackExpiry",
    "track_serial": "TrackSerial",
    "allow_negative_stock": "AllowNegativeStock",
    "remarks": "Remarks",
}

_TEXT_FIELDS: dict[str, str] = {
    "Brand": "brand",
    "Model": "model",
    "Barcode": "barcode",
    "ShortName": "short_name",
    "Description": "description",
    "Remarks": "remarks",
}
_MONEY_FIELDS: dict[str, str] = {
    "PurchasePrice": "purchase_price",
    "SellingPrice": "selling_price",
    "MRP": "mrp",
}
_FLAG_FIELDS: dict[str, str] = {
    "TrackBatch": "track_batch",
    "TrackExpiry": "track_expiry",
    "TrackSerial": "track_serial",
    "AllowNegativeStock": "allow_negative_stock",
}
_YES = {"yes", "y", "true", "1", "t"}
_NO = {"no", "n", "false", "0", "f"}


def normalise_heading(value: object) -> str:
    """Reduce a heading to letters and digits, lower case."""
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


_HEADING_LOOKUP: dict[str, str] = {}
for _column in COLUMNS:
    _HEADING_LOOKUP[normalise_heading(_column.heading)] = _column.heading
    for _alias in _column.aliases:
        _HEADING_LOOKUP[_alias] = _column.heading


@dataclass
class ImportIssue:
    """One problem with one row: where, and what."""

    row: int
    code: str | None
    column: str | None
    message: str

    def describe(self) -> str:
        """Render it the way every import error in this module reads."""
        where = f"Row {self.row}" + (f" ({self.code})" if self.code else "")
        column = f"{self.column}: " if self.column else ""
        return f"{where}: {column}{self.message}"


@dataclass
class ImportReport:
    """What a file would do, or did."""

    rows: int = 0
    to_create: int = 0
    to_update: int = 0
    skipped_blank: int = 0
    columns_used: list[str] = field(default_factory=list)
    columns_ignored: list[str] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    imported: bool = False
    products: list[Product] = field(default_factory=list)


@dataclass
class _Row:
    """One data row read off the file, by canonical heading."""

    number: int
    cells: dict[str, str]


class _References:
    """The firm's categories, units and tax groups, read once per file."""

    def __init__(self, session: Session, firm_id: UUID) -> None:
        categories = session.scalars(
            select(ProductCategory).where(
                ProductCategory.firm_id == firm_id,
                ProductCategory.is_deleted.is_(False),
            )
        ).all()
        self.categories = list(categories)
        units = session.scalars(
            select(Uom).where(Uom.is_deleted.is_(False), Uom.status == "ACTIVE")
        ).all()
        self.units: dict[str, Uom] = {}
        # Codes win over names and symbols: filled last so they overwrite.
        for unit in units:
            if unit.symbol:
                self.units.setdefault(unit.symbol.strip().lower(), unit)
        for unit in units:
            self.units[unit.name.strip().lower()] = unit
        for unit in units:
            self.units[unit.code.strip().lower()] = unit
        self.tax_groups = {
            code.upper(): code
            for code in session.scalars(
                select(TaxProfile.group_code).where(
                    TaxProfile.firm_id == firm_id,
                    TaxProfile.is_deleted.is_(False),
                    TaxProfile.status == "ACTIVE",
                )
            ).all()
            if code
        }

    def category(self, value: str, parent_id: UUID | None) -> ProductCategory | None:
        """Find a category by code, then by name, under ``parent_id``.

        A top-level lookup (``parent_id`` None) accepts any category, so a
        firm whose masters are one level deep is not refused.
        """
        wanted = value.strip().lower()
        pool = [
            item
            for item in self.categories
            if parent_id is None or item.parent_id == parent_id
        ]
        by_code = [item for item in pool if item.code.lower() == wanted]
        by_name = [item for item in pool if item.name.strip().lower() == wanted]
        for found in (by_code, by_name):
            if found:
                # Prefer the shallowest: a top-level category of that name
                # over a sub category somewhere else that shares it.
                return min(found, key=lambda item: item.level)
        return None


def read_rows(
    content: bytes, file_format: Literal["csv", "xlsx"]
) -> tuple[list[_Row], list[str], list[str]]:
    """Read a file into rows keyed by canonical heading.

    Returns the rows, the headings used and the headings ignored. The header
    is row 1, as a spreadsheet shows it, so data starts at row 2.
    """
    grid = _read_csv(content) if file_format == "csv" else _read_xlsx(content)
    try:
        header = next(grid)
    except StopIteration:
        return [], [], []
    positions: dict[str, int] = {}
    ignored: list[str] = []
    for position, raw in enumerate(header):
        heading = _HEADING_LOOKUP.get(normalise_heading(raw))
        if heading is None or heading in positions:
            if str(raw or "").strip():
                ignored.append(str(raw).strip())
            continue
        positions[heading] = position
    rows: list[_Row] = []
    for number, values in enumerate(grid, start=2):
        cells = {
            heading: _cell_text(values[position]) if position < len(values) else ""
            for heading, position in positions.items()
        }
        rows.append(_Row(number=number, cells=cells))
    used = [column.heading for column in COLUMNS if column.heading in positions]
    return rows, used, ignored


def _read_csv(content: bytes) -> Iterator[list[object]]:
    """Yield CSV rows, taking the byte-order mark Excel writes."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows saves "CSV" in the ANSI code page.
        text = content.decode("cp1252")
    yield from csv.reader(io.StringIO(text))


def _read_xlsx(content: bytes) -> Iterator[list[object]]:
    """Yield the rows of the workbook's first sheet."""
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as error:
        raise ValidationError(
            "XLSX import dependency is unavailable. Install openpyxl."
        ) from error
    try:
        workbook = load_workbook(filename=BytesIO(content), read_only=True)
    except Exception as error:  # noqa: BLE001 -- any unreadable file
        raise ValidationError("The file is not a readable XLSX workbook.") from error
    sheet = workbook.worksheets[0]
    for values in sheet.iter_rows(values_only=True):
        yield list(values)


def _cell_text(value: object) -> str:
    """Read a cell as trimmed text; a whole float loses its ``.0``."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class ProductFileImporter:
    """Check a product file row by row, and import it whole or not at all."""

    def __init__(self, session: Session, service: "ProductService") -> None:
        self._session = session
        self._service = service

    def run(
        self,
        content: bytes,
        *,
        file_format: Literal["csv", "xlsx"],
        firm_id: UUID,
        actor_id: UUID,
        existing: ExistingRows = "refuse",
        apply: bool,
    ) -> ImportReport:
        """Check every row; with ``apply``, commit the file if nothing failed.

        Nothing is written by a check, and nothing is written by an apply
        that found a problem: the staged rows are rolled back either way.
        """
        rows, used, ignored = read_rows(content, file_format)
        report = ImportReport(columns_used=used, columns_ignored=ignored)
        missing = [
            column.heading
            for column in COLUMNS
            if column.required and column.heading not in used
        ]
        if missing:
            report.issues.append(
                ImportIssue(
                    row=1,
                    code=None,
                    column=None,
                    message="The heading row has no "
                    + " or ".join(missing)
                    + " column. Download the template to see the headings.",
                )
            )
            return report
        references = _References(self._session, firm_id)
        stored = {
            product.code: product
            for product in self._session.scalars(
                select(Product).where(
                    Product.firm_id == firm_id, Product.is_deleted.is_(False)
                )
            ).all()
        }
        seen: dict[str, int] = {}
        try:
            for row in rows:
                if not any(row.cells.values()):
                    report.skipped_blank += 1
                    continue
                report.rows += 1
                code = row.cells.get("Code", "").strip().upper()
                if not code:
                    report.issues.append(
                        ImportIssue(row.number, None, "Code", "is required.")
                    )
                    continue
                if code in seen:
                    report.issues.append(
                        ImportIssue(
                            row.number,
                            code,
                            "Code",
                            f"appears again; row {seen[code]} already has it.",
                        )
                    )
                    continue
                seen[code] = row.number
                current = stored.get(code)
                if current is not None and existing == "refuse":
                    report.issues.append(
                        ImportIssue(
                            row.number,
                            code,
                            "Code",
                            "is already a product. Choose to update existing "
                            "products to change it from this file.",
                        )
                    )
                    continue
                self._stage(row, code, current, references, firm_id, actor_id, report)
        except Exception:
            self._session.rollback()
            raise
        if report.rows == 0 and not report.issues:
            report.issues.append(
                ImportIssue(1, None, None, "The file has no product rows.")
            )
        if not apply or report.issues:
            self._session.rollback()
            report.products = []
            return report
        self._service._commit()
        for product in report.products:
            self._session.refresh(product)
        report.imported = True
        return report

    def _stage(
        self,
        row: _Row,
        code: str,
        current: Product | None,
        references: _References,
        firm_id: UUID,
        actor_id: UUID,
        report: ImportReport,
    ) -> None:
        """Build, validate and stage one row; record what went wrong if not."""
        issues: list[ImportIssue] = []
        values = self._values(row, code, current, references, issues)
        if issues:
            report.issues.extend(issues)
            return
        try:
            data = (
                ProductCreate.model_validate(values)
                if current is None
                else ProductUpdate.model_validate(values)
            )
        except PydanticValidationError as error:
            for detail in error.errors():
                name = str(detail["loc"][0]) if detail["loc"] else ""
                message = str(detail["msg"]).removeprefix("Value error, ")
                report.issues.append(
                    ImportIssue(
                        row.number, code, _FIELD_HEADINGS.get(name), message + "."
                    )
                )
            return
        try:
            if isinstance(data, ProductCreate):
                product = self._service.stage_product(
                    data, firm_id=firm_id, actor_id=actor_id
                )
                report.to_create += 1
            elif current is not None:
                product = self._service.stage_update_product(
                    current, data, firm_scope=firm_id, actor_id=actor_id
                )
                report.to_update += 1
        except ApplicationError as error:
            # The guards all run before the row is added, so the session is
            # still sound and the rest of the file can be checked.
            report.issues.append(ImportIssue(row.number, code, None, error.message))
            return
        report.products.append(product)

    def _values(
        self,
        row: _Row,
        code: str,
        current: Product | None,
        references: _References,
        issues: list[ImportIssue],
    ) -> dict[str, object]:
        """Turn one row's cells into write-schema values.

        On a create a blank cell takes the schema default. On an update a blank
        cell -- or a column the file does not have -- is left out, so it leaves
        the stored value alone.
        """
        cells = row.cells

        def fail(column: str, message: str) -> None:
            """Record one problem with this row."""
            issues.append(ImportIssue(row.number, code, column, message))

        values: dict[str, object] = {"code": code}
        name = cells.get("Name", "")
        if name:
            values["name"] = name
        elif current is None:
            fail("Name", "is required.")
        else:
            values["name"] = current.name
        kind = cells.get("Type", "").upper().replace(" ", "_")
        if kind:
            values["product_type"] = kind
        else:
            values["product_type"] = (
                current.product_type if current is not None else "STOCK_ITEM"
            )
        status = cells.get("Status", "").upper()
        if status:
            values["status"] = status
        for heading, target in _TEXT_FIELDS.items():
            if cells.get(heading):
                values[target] = cells[heading]
        hsn = cells.get("HSN", "")
        if hsn:
            values["hsn_sac"] = hsn.upper()
        for heading, target in _MONEY_FIELDS.items():
            raw = cells.get(heading, "")
            if not raw:
                continue
            try:
                values[target] = Decimal(raw.replace(",", "").strip())
            except InvalidOperation:
                fail(heading, f"'{raw}' is not a number.")
        for heading, target in _FLAG_FIELDS.items():
            raw = cells.get(heading, "").lower()
            if not raw:
                continue
            if raw in _YES:
                values[target] = True
            elif raw in _NO:
                values[target] = False
            else:
                fail(heading, f"'{cells[heading]}' should be Yes or No.")
        self._resolve(cells, current, references, values, fail)
        return values

    @staticmethod
    def _resolve(
        cells: dict[str, str],
        current: Product | None,
        references: _References,
        values: dict[str, object],
        fail: Callable[[str, str], None],
    ) -> None:
        """Turn category, unit and tax group names into what the schema takes."""
        category_id: UUID | None = current.category_id if current else None
        category_text = cells.get("Category", "")
        if category_text:
            category = references.category(category_text, None)
            if category is None:
                fail("Category", f"'{category_text}' is not one of the firm's.")
            else:
                category_id = category.id
                values["category_id"] = category.id
        sub_text = cells.get("SubCategory", "")
        if sub_text:
            if category_id is None:
                fail("SubCategory", "needs a Category.")
            else:
                sub = references.category(sub_text, category_id)
                if sub is None:
                    fail(
                        "SubCategory",
                        f"'{sub_text}' is not a sub category of that category.",
                    )
                else:
                    values["sub_category_id"] = sub.id
        elif (
            category_text
            and current is not None
            and category_id != (current.category_id)
        ):
            # A new category cannot keep the old one's sub category.
            values["sub_category_id"] = None
        unit_text = cells.get("Unit", "")
        if unit_text:
            unit = references.units.get(unit_text.strip().lower())
            if unit is None:
                fail("Unit", f"'{unit_text}' is not a unit on this system.")
            else:
                values["unit"] = unit.code
                values["base_uom_id"] = unit.id
                # A new product is one unit throughout until packaging is set
                # up; an existing one keeps any other units it already has.
                for target in ("inventory_uom_id", "purchase_uom_id", "sales_uom_id"):
                    if current is None or getattr(current, target) is None:
                        values[target] = unit.id
        tax_text = cells.get("TaxGroup", "")
        if tax_text:
            group = references.tax_groups.get(tax_text.strip().upper())
            if group is None:
                fail("TaxGroup", f"'{tax_text}' is not an active tax group.")
            else:
                values["tax_profile_group_code"] = group


def template_workbook(session: Session, firm_id: UUID) -> bytes:
    """Build the XLSX template: the sheet to fill, the notes, and the lists."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font  # type: ignore[import-untyped]
    except ImportError as error:
        raise ValidationError(
            "XLSX export dependency is unavailable. Install openpyxl."
        ) from error
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append([column.heading for column in COLUMNS])
    sheet.append([column.example for column in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, column in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = max(12, len(column.heading) + 4)
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("Notes")
    notes.append(["Column", "Required", "What it takes", "Also read from"])
    for column in COLUMNS:
        notes.append(
            [
                column.heading,
                "Yes" if column.required else "No",
                column.takes,
                ", ".join(column.aliases),
            ]
        )
    notes.append([])
    for line in (
        "Replace the example row with your products; only the Products sheet "
        "is read.",
        "Headings are matched ignoring case and spaces, and the names in "
        "'Also read from' are accepted too; other columns are ignored.",
        "Nothing is saved until every row passes the check.",
        "When updating existing products, a blank cell leaves that field " "as it is.",
        "Opening stock is imported separately, under Inventory.",
    ):
        notes.append([line])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["C"].width = 70
    notes.column_dimensions["D"].width = 50

    references = _References(session, firm_id)
    lists = workbook.create_sheet("Lists")
    lists.append(
        [
            "Category code",
            "Category name",
            "Parent",
            "",
            "Unit code",
            "Unit name",
            "",
            "Tax group",
        ]
    )
    parents = {item.id: item.code for item in references.categories}
    categories = sorted(references.categories, key=lambda item: item.path)
    units = sorted(
        {unit.id: unit for unit in references.units.values()}.values(),
        key=lambda unit: unit.code,
    )
    groups = sorted(references.tax_groups.values())
    for index in range(max(len(categories), len(units), len(groups))):
        category = categories[index] if index < len(categories) else None
        unit = units[index] if index < len(units) else None
        group = groups[index] if index < len(groups) else None
        lists.append(
            [
                category.code if category else None,
                category.name if category else None,
                (
                    parents.get(category.parent_id)
                    if category and category.parent_id
                    else None
                ),
                None,
                unit.code if unit else None,
                unit.name if unit else None,
                None,
                group,
            ]
        )
    for cell in lists[1]:
        cell.font = Font(bold=True)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template_csv() -> str:
    """Build the CSV template: the headings and one example row."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([column.heading for column in COLUMNS])
    writer.writerow([column.example for column in COLUMNS])
    return buffer.getvalue()
