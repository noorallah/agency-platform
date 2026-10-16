"""Brands come in from a file, with a template to fill in (Swathi G3).

The file rules are the other masters', from ``app.common.file_import``. What
these tests pin is what is particular to brands -- the name is the key, a new
brand names an existing company, and every row goes through the service's
own guards before the file commits once.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.firms.models import Firm
from app.imports.services import IMPORT_KINDS, columns_for_kind
from app.products.models.brand import Brand, Principal
from app.products.services.brand_import import (
    COLUMNS,
    BrandFileImporter,
    template_csv,
    template_workbook,
)
from app.products.services.brands import BrandService


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with two companies to file brands under."""
    firm = Firm(
        name="Brand Firm",
        code="BRANDS",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    session.add_all(
        [
            Principal(firm_id=firm.id, code="HUL", name="Hindustan Unilever"),
            Principal(firm_id=firm.id, code="ITC", name="ITC Limited"),
        ]
    )
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
) -> ImportReport[Brand]:
    """Check or apply one file as a single actor."""
    return BrandFileImporter(session, BrandService(session)).run(
        content,
        file_format="csv",
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _names(factory: sessionmaker[Session]) -> list[str]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(
            fresh.scalars(select(Brand.name).where(Brand.is_deleted.is_(False))).all()
        )


def test_the_template_names_the_columns_and_lists_the_firms_companies() -> None:
    """The CSV carries the headings; the workbook lists the companies."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        assert template_csv().splitlines()[0] == ",".join(c.heading for c in COLUMNS)
        workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))
        assert workbook.sheetnames == ["Brands", "Notes", "Lists"]
        lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
        assert lists[0] == ["Company code", "Company name"]
        assert lists[1:] == [
            ["HUL", "Hindustan Unilever"],
            ["ITC", "ITC Limited"],
        ]


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        report = _run(
            session,
            firm.id,
            _csv(
                "Brand Name,Company,Active",
                "Surf,HUL,Yes",
                "Lux,,Yes",
                "Classmate,NOBODY,",
                "Bingo,ITC,maybe",
                "surf,HUL,",
                ",HUL,",
                "Sunfeast,ITC Limited,No",
            ),
            apply=False,
        )
    found = {(issue.row, issue.column, issue.message) for issue in report.issues}
    assert (3, "Company", "is needed for a new brand.") in found
    assert any(row == 4 and column == "Company" for row, column, _ in found)
    assert (5, "Active", "'maybe' should be Yes or No.") in found
    assert (6, "Name", "appears again; row 2 already has it.") in found
    assert (7, "Name", "is required.") in found
    assert len(report.issues) == 5
    assert report.to_create == 2  # Surf and Sunfeast staged, then rolled back
    assert report.imported is False
    assert _names(factory) == []


def test_an_apply_is_all_or_nothing() -> None:
    """One bad row keeps the good ones out; a clean file goes in whole."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        bad = _run(
            session,
            firm.id,
            _csv("Name,Company", "Surf,HUL", "Lux,NOBODY"),
            apply=True,
        )
        assert bad.imported is False
        assert _names(factory) == []

        good = _run(
            session,
            firm.id,
            _csv("Name,Company,Active", "Surf,HUL,", "Sunfeast,ITC Limited,No"),
            apply=True,
        )
    assert good.imported is True
    assert good.to_create == 2
    assert _names(factory) == ["Sunfeast", "Surf"]
    with factory() as fresh:
        sunfeast = fresh.scalars(select(Brand).where(Brand.name == "Sunfeast")).one()
        itc = fresh.scalars(select(Principal).where(Principal.code == "ITC")).one()
        assert sunfeast.principal_id == itc.id
        assert sunfeast.is_active is False
        # Each brand is audited as the form audits it.
        assert (
            len(
                fresh.scalars(
                    select(AuditLog).where(AuditLog.action == "brand.saved")
                ).all()
            )
            == 2
        )


def test_a_second_apply_refuses_the_same_names_unless_asked_to_update() -> None:
    """A name already a brand is refused, ignoring case; update changes it."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        first = _run(session, firm.id, _csv("Name,Company", "Surf,HUL"), apply=True)
        assert first.imported is True

        again = _run(session, firm.id, _csv("Name,Company", "SURF,ITC"), apply=True)
        assert again.imported is False
        assert [(i.row, i.column) for i in again.issues] == [(2, "Name")]
        assert "is already a brand" in again.issues[0].message

        updated = _run(
            session,
            firm.id,
            _csv("Name,Company,Active", "SURF,,No"),
            apply=True,
            existing="update",
        )
    assert updated.imported is True
    assert updated.to_update == 1
    with factory() as fresh:
        brands = fresh.scalars(select(Brand)).all()
        assert len(brands) == 1
        hul = fresh.scalars(select(Principal).where(Principal.code == "HUL")).one()
        # A blank Company keeps the one it had; the name takes the file's case.
        assert brands[0].principal_id == hul.id
        assert brands[0].name == "SURF"
        assert brands[0].is_active is False


def test_brands_can_be_mapped_from_another_softwares_file() -> None:
    """The kind is registered under the brand code, with the brand columns."""
    factory = _factory()
    assert IMPORT_KINDS["brands"].permission == "PRODUCT_UPDATE"
    with factory() as session:
        assert columns_for_kind(session, "brands") == COLUMNS
