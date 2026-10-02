"""PAN and TAN checked on save, and the quarterly TDS return file (53, 53.1).

Backlog 53 item 2: a PAN is ``AAAAA9999A``, a TAN ``AAAA99999A``, and
characters 3 to 12 of a GSTIN are its holder's PAN -- so a blank PAN is
filled from the GSTIN and one that disagrees is refused naming both. A value
typed before any check existed must not block an unrelated edit, so only what
a write *sets* is checked.

Backlog 53.1: the 26Q export -- one quarter's deductions laid out as the
return lists them, as a workbook to prepare the return from.
"""

# ruff: noqa: D103

from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.file_import import ImportReport
from app.common.scope import ResolvedFirmScope
from app.core.database.base import Base
from app.core.enums import TokenType
from app.core.exceptions import ValidationError
from app.core.security.authorization import Principal
from app.core.security.jwt import TokenClaims
from app.core.validation import check_tan_if_set, pan_in_gstin, settle_pan
from app.customers.models import Customer
from app.customers.schemas import CustomerCreate, CustomerUpdate
from app.customers.services import CustomerService
from app.customers.services.customer_import import CustomerFileImporter
from app.expenses.schemas import ExpenseCreate
from app.expenses.services import ExpenseService
from app.finance.api.router import tds_26q_deductees, tds_26q_file
from app.finance.models import LedgerAccount
from app.finance.services.tds_return import (
    DEDUCTEE_HEADINGS,
    PAN_NOT_AVAILABLE,
    TdsReturnService,
    deductees_csv,
    due_date,
    return_quarter,
    return_workbook,
)
from app.firms.models import Firm
from app.firms.schemas import FirmCreate
from app.firms.services import FirmService
from app.settlements.schemas import (
    SettlementAllocationWrite,
    SettlementCreate,
    SettlementMethodEnum,
)
from app.settlements.services import PaymentService
from app.vendors.models import Vendor
from app.vendors.schemas import VendorCreate, VendorUpdate
from app.vendors.services import VendorService
from app.vendors.services.vendor_import import VendorFileImporter
from tests.unit.test_settlements import _Books, _session_factory

pytestmark = pytest.mark.typed_document_numbers

ACTOR = uuid4()
COMPANY_PAN = "AAACP1234C"
COMPANY_GSTIN = f"29{COMPANY_PAN}1Z5"


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    firm = Firm(
        name="Pan Firm",
        code="PANFIRM",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.commit()
    return firm


def _customer(code: str, **fields: object) -> CustomerCreate:
    return CustomerCreate.model_validate(
        {
            "code": code,
            "customer_type": "BUSINESS",
            "name": f"Customer {code}",
            "currency_code": "INR",
            **fields,
        }
    )


def _vendor(code: str, **fields: object) -> VendorCreate:
    return VendorCreate.model_validate(
        {"code": code, "name": f"Vendor {code}", **fields}
    )


# --------------------------------------------------------------------------
# The rules themselves
# --------------------------------------------------------------------------


def test_a_gstin_carries_its_holders_pan() -> None:
    assert pan_in_gstin(COMPANY_GSTIN) == COMPANY_PAN
    assert pan_in_gstin(COMPANY_GSTIN.lower()) == COMPANY_PAN
    # A registration not built on a PAN says nothing about one.
    assert pan_in_gstin("0717UNO00157UN5") is None
    assert pan_in_gstin("GST-001") is None
    assert pan_in_gstin(None) is None


def test_a_malformed_pan_is_refused_naming_the_field() -> None:
    with pytest.raises(ValidationError, match="five letters, four digits") as error:
        settle_pan(
            pan="ABCD1234F",
            gstin=None,
            creating=True,
            pan_field="pan_number",
            gstin_field="gst_number",
        )
    assert error.value.details == {"field": "pan_number"}


def test_a_blank_pan_is_filled_from_the_gstin_and_a_different_one_refused() -> None:
    filled = settle_pan(
        pan=None,
        gstin=COMPANY_GSTIN,
        creating=True,
        pan_field="pan",
        gstin_field="gstin",
    )
    assert filled == COMPANY_PAN
    with pytest.raises(ValidationError) as error:
        settle_pan(
            pan="AAACQ9999Z",
            gstin=COMPANY_GSTIN,
            creating=True,
            pan_field="pan",
            gstin_field="gstin",
        )
    assert "AAACQ9999Z" in error.value.message
    assert COMPANY_GSTIN in error.value.message
    assert error.value.details == {"field": "pan", "fields": ["pan", "gstin"]}


def test_a_stored_bad_pan_is_left_alone_until_somebody_sets_it() -> None:
    # Resent unchanged: not being set, so not checked.
    assert (
        settle_pan(
            pan="PAN-OLD",
            gstin=None,
            stored_pan="PAN-OLD",
            creating=False,
            pan_field="pan",
            gstin_field="gstin",
        )
        == "PAN-OLD"
    )
    with pytest.raises(ValidationError):
        settle_pan(
            pan="PAN-NEW",
            gstin=None,
            stored_pan="PAN-OLD",
            creating=False,
            pan_field="pan",
            gstin_field="gstin",
        )


def test_a_tan_is_checked_only_when_set() -> None:
    assert check_tan_if_set("dela12345b", creating=True, field="tan") == "DELA12345B"
    assert (
        check_tan_if_set("OLD-TAN", stored_tan="OLD-TAN", creating=False, field="tan")
        == "OLD-TAN"
    )
    with pytest.raises(ValidationError, match="four letters, five digits") as error:
        check_tan_if_set("DEL12345B", creating=True, field="tax.tan")
    assert error.value.details == {"field": "tax.tan"}


# --------------------------------------------------------------------------
# Customers
# --------------------------------------------------------------------------


def test_a_customer_pan_is_checked_filled_and_matched() -> None:
    session = _factory()()
    firm = _firm(session)
    service = CustomerService(session)

    with pytest.raises(ValidationError, match="PAN AB12 is not a PAN"):
        service.create(
            _customer("BAD", pan_number="ab12"), firm_id=firm.id, actor_id=ACTOR
        )
    with pytest.raises(ValidationError, match="does not match GSTIN"):
        service.create(
            _customer("CLASH", gst_number=COMPANY_GSTIN, pan_number="AAACQ9999Z"),
            firm_id=firm.id,
            actor_id=ACTOR,
        )
    filled = service.create(
        _customer("GOOD", gst_number=COMPANY_GSTIN), firm_id=firm.id, actor_id=ACTOR
    )
    assert filled.pan_number == COMPANY_PAN


def test_a_second_branch_of_the_same_company_keeps_its_pan_and_is_named() -> None:
    """One company, a GSTIN per state, one PAN on both (decision A7).

    Until A7 the second branch's PAN was left blank, because a PAN was unique
    among customers. Now it is kept, and the save is warned about by name.
    """
    session = _factory()()
    firm = _firm(session)
    service = CustomerService(session)
    first = service.create(
        _customer("KA", gst_number=COMPANY_GSTIN), firm_id=firm.id, actor_id=ACTOR
    )
    other_state = f"33{COMPANY_PAN}1Z9"
    second = service.create(
        _customer("TN", gst_number=other_state), firm_id=firm.id, actor_id=ACTOR
    )
    assert second.pan_number == COMPANY_PAN
    warning = service.identity_warning(
        firm.id,
        gst_number=second.gst_number,
        pan_number=second.pan_number,
        excluding_id=second.id,
    )
    assert warning == f"PAN {COMPANY_PAN} is also on {first.code} {first.name}."


def test_a_customer_with_an_old_bad_pan_can_still_be_edited() -> None:
    session = _factory()()
    firm = _firm(session)
    service = CustomerService(session)
    customer = service.create(_customer("OLD"), firm_id=firm.id, actor_id=ACTOR)
    customer.pan_number = "PAN-OLD"  # typed before the check existed
    session.commit()

    # A partial edit, and a complete form resending the stored PAN, both pass.
    service.update(
        customer.id,
        CustomerUpdate.model_validate(
            {
                "code": "OLD",
                "customer_type": "BUSINESS",
                "name": "Renamed",
                "currency_code": "INR",
            }
        ),
        firm_scope=firm.id,
        actor_id=ACTOR,
    )
    service.update(
        customer.id,
        CustomerUpdate.model_validate(
            {
                "code": "OLD",
                "customer_type": "BUSINESS",
                "name": "Renamed again",
                "currency_code": "INR",
                "pan_number": "PAN-OLD",
                "phone": "+919876543210",
            }
        ),
        firm_scope=firm.id,
        actor_id=ACTOR,
    )
    assert customer.name == "Renamed again"
    assert customer.pan_number == "PAN-OLD"

    # Setting the PAN is when it is checked.
    with pytest.raises(ValidationError, match="is not a PAN"):
        service.update(
            customer.id,
            CustomerUpdate.model_validate(
                {
                    "code": "OLD",
                    "customer_type": "BUSINESS",
                    "name": "Renamed again",
                    "currency_code": "INR",
                    "pan_number": "PAN-NEW",
                }
            ),
            firm_scope=firm.id,
            actor_id=ACTOR,
        )


def test_a_gstin_added_later_fills_a_blank_pan() -> None:
    session = _factory()()
    firm = _firm(session)
    service = CustomerService(session)
    customer = service.create(_customer("LATER"), firm_id=firm.id, actor_id=ACTOR)
    service.update(
        customer.id,
        CustomerUpdate.model_validate(
            {
                "code": "LATER",
                "customer_type": "BUSINESS",
                "name": "Customer LATER",
                "currency_code": "INR",
                "gst_number": COMPANY_GSTIN,
            }
        ),
        firm_scope=firm.id,
        actor_id=ACTOR,
    )
    assert customer.pan_number == COMPANY_PAN


def test_a_customer_file_names_the_pan_column_on_the_rows_that_fail() -> None:
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = (
        "Code,Name,GSTIN,PAN\n"
        f"OK,Fine,{COMPANY_GSTIN},\n"
        "SHAPE,Bad shape,,ABC123\n"
        f"CLASH,Mismatch,33{COMPANY_PAN}1Z9,AAACQ9999Z\n"
    ).encode()
    report: ImportReport[Customer] = CustomerFileImporter(
        session, CustomerService(session), may_manage_settings=True
    ).run(
        content,
        file_format="csv",
        firm_id=firm.id,
        actor_id=ACTOR,
        existing="refuse",
        apply=False,
    )
    found = {(issue.row, issue.code, issue.column) for issue in report.issues}
    assert found == {(3, "SHAPE", "PAN"), (4, "CLASH", "PAN")}, [
        issue.describe() for issue in report.issues
    ]
    assert "AAACQ9999Z" in report.issues[-1].message


# --------------------------------------------------------------------------
# Vendors and the firm
# --------------------------------------------------------------------------


def test_a_vendor_pan_and_tax_row_tan_are_checked_when_set() -> None:
    session = _factory()()
    firm = _firm(session)
    service = VendorService(session)

    with pytest.raises(ValidationError, match="is not a PAN"):
        service.create(_vendor("V-BAD", pan="12345"), firm_id=firm.id, actor_id=ACTOR)
    with pytest.raises(ValidationError, match="is not a TAN") as error:
        service.create(
            _vendor("V-TAN", tax=[{"tan": "DEL12345B", "is_primary": True}]),
            firm_id=firm.id,
            actor_id=ACTOR,
        )
    assert error.value.details == {"field": "tax.tan"}

    vendor = service.create(
        _vendor(
            "V-OK",
            gstin=COMPANY_GSTIN,
            tax=[{"gstin": COMPANY_GSTIN, "tan": "blra12345c", "is_primary": True}],
        ),
        firm_id=firm.id,
        actor_id=ACTOR,
    )
    assert vendor.pan == COMPANY_PAN
    [row] = vendor.tax_details
    assert (row.pan, row.tan) == (COMPANY_PAN, "BLRA12345C")


def test_a_vendor_with_an_old_bad_tan_can_still_be_edited() -> None:
    session = _factory()()
    firm = _firm(session)
    service = VendorService(session)
    vendor = service.create(
        _vendor("V-OLD", tax=[{"is_primary": True}]), firm_id=firm.id, actor_id=ACTOR
    )
    [row] = vendor.tax_details
    row.tan = "OLD-TAN"
    vendor.pan = "PAN-OLD"
    session.commit()

    # The desktop resends every collection it shows, unchanged.
    service.update(
        vendor.id,
        VendorUpdate.model_validate(
            {
                "code": "V-OLD",
                "name": "Renamed",
                "pan": "PAN-OLD",
                "tax": [{"id": str(row.id), "tan": "OLD-TAN", "is_primary": True}],
            }
        ),
        firm_scope=firm.id,
        actor_id=ACTOR,
    )
    assert vendor.name == "Renamed"
    assert session.scalar(select(Vendor.pan).where(Vendor.id == vendor.id)) == (
        "PAN-OLD"
    )
    with pytest.raises(ValidationError, match="is not a TAN"):
        service.update(
            vendor.id,
            VendorUpdate.model_validate(
                {
                    "code": "V-OLD",
                    "name": "Renamed",
                    "tax": [{"id": str(row.id), "tan": "NEW-TAN", "is_primary": True}],
                }
            ),
            firm_scope=firm.id,
            actor_id=ACTOR,
        )


def test_a_vendor_file_names_the_pan_column() -> None:
    factory = _factory()
    session = factory()
    firm = _firm(session)
    report = VendorFileImporter(
        session, VendorService(session), may_manage_bank_details=True
    ).run(
        b"Code,Name,PAN\nV-1,Bad,ABC123\n",
        file_format="csv",
        firm_id=firm.id,
        actor_id=ACTOR,
        existing="refuse",
        apply=False,
    )
    assert [(issue.row, issue.column) for issue in report.issues] == [(2, "PAN")]


def test_the_firms_own_pan_must_match_its_gstin() -> None:
    session = _factory()()
    service = FirmService(session)
    payload = {
        "name": "Firm",
        "code": "FIRMX",
        "country": "IN",
        "currency_code": "INR",
        "financial_year_start": date(2026, 4, 1),
    }
    with pytest.raises(ValidationError, match="does not match GSTIN"):
        service.create(
            FirmCreate.model_validate(
                {**payload, "gst_number": COMPANY_GSTIN, "pan_number": "AAACQ9999Z"}
            ),
            ACTOR,
        )
    firm = service.create(
        FirmCreate.model_validate({**payload, "gst_number": COMPANY_GSTIN}), ACTOR
    )
    assert firm.pan_number == COMPANY_PAN


# --------------------------------------------------------------------------
# The 26Q export
# --------------------------------------------------------------------------


def test_a_return_is_named_by_year_and_quarter() -> None:
    first = return_quarter("2026-27", "Q1")
    assert (first.start, first.end) == (date(2026, 4, 1), date(2026, 6, 30))
    third = return_quarter("2026-2027", "3")
    assert (third.start, third.end) == (date(2026, 10, 1), date(2026, 12, 31))
    last = return_quarter("2026-27", "q4")
    assert (last.start, last.end) == (date(2027, 1, 1), date(2027, 3, 31))
    assert last.assessment_year == "2027-28"
    assert last.file_stem == "26Q-2026-27-Q4"
    for year, quarter in (("2026", "Q1"), ("2026-28", "Q1"), ("2026-27", "Q5")):
        with pytest.raises(ValidationError):
            return_quarter(year, quarter)


def test_tax_deducted_is_due_on_the_seventh_and_march_on_the_thirtieth() -> None:
    assert due_date(date(2026, 4, 20)) == date(2026, 5, 7)
    assert due_date(date(2026, 12, 3)) == date(2027, 1, 7)
    assert due_date(date(2027, 3, 31)) == date(2027, 4, 30)


def _account(books: _Books, code: str) -> UUID:
    account_id = books.session.scalar(
        select(LedgerAccount.id).where(
            LedgerAccount.firm_id == books.firm.id, LedgerAccount.code == code
        )
    )
    assert account_id is not None
    return account_id


def _expense(
    books: _Books,
    *,
    when: date,
    amount: str,
    tds: str,
    section: str,
    payee: str,
    pan: str | None = None,
) -> None:
    ExpenseService(books.session).create(
        ExpenseCreate(
            expense_date=when,
            expense_account_id=_account(books, "6000"),
            paid_from_account_id=_account(books, "1010"),
            amount=Decimal(amount),
            payee=payee,
            payee_pan=pan,
            tds_amount=Decimal(tds),
            tds_section=section,
        ),
        firm_id=books.firm.id,
        actor_id=books.actor_id,
    )
    books.session.commit()


def _seeded_quarter() -> _Books:
    """Seed Q1 2026-27: a 194Q payment, a fee to a payee with no PAN.

    Beside them a payment later reversed, a salary deduction, and one
    deduction in Q2: two the return leaves out, one outside its quarter.
    """
    books = _Books(_session_factory()())
    books.firm.tan_number = "BLRA12345C"
    books.firm.pan_number = "AAAFA1234B"
    books.vendor.pan = COMPANY_PAN
    books.session.commit()
    invoice = books.purchase_invoice("PI-1", "100000.00")
    payments = PaymentService(books.session)
    payments.create(
        SettlementCreate(
            party_id=books.vendor.id,
            settlement_date=date(2026, 4, 20),
            amount=Decimal("100000.00"),
            method=SettlementMethodEnum.BANK,
            tds_amount=Decimal("100.00"),
            tds_section="194Q",
            allocations=[
                SettlementAllocationWrite(
                    invoice_id=invoice.id, amount=Decimal("100000.00")
                )
            ],
        ),
        firm_id=books.firm.id,
        actor_id=books.actor_id,
    )
    books.session.commit()
    second = books.purchase_invoice("PI-2", "5000.00")
    reversed_payment = payments.create(
        SettlementCreate(
            party_id=books.vendor.id,
            settlement_date=date(2026, 4, 20),
            amount=Decimal("5000.00"),
            method=SettlementMethodEnum.BANK,
            tds_amount=Decimal("50.00"),
            tds_section="194C",
            allocations=[
                SettlementAllocationWrite(
                    invoice_id=second.id, amount=Decimal("5000.00")
                )
            ],
        ),
        firm_id=books.firm.id,
        actor_id=books.actor_id,
    )
    books.session.commit()
    payments.reverse(
        reversed_payment.id, firm_id=books.firm.id, actor_id=books.actor_id
    )
    books.session.commit()
    _expense(
        books,
        when=date(2026, 5, 5),
        amount="30000.00",
        tds="6000.00",
        section="194J",
        payee="Dr Rao",
    )
    _expense(
        books,
        when=date(2026, 6, 1),
        amount="50000.00",
        tds="5000.00",
        section="192",
        payee="A Clerk",
        pan="ABCPC1234D",
    )
    _expense(
        books,
        when=date(2026, 7, 1),
        amount="10000.00",
        tds="1000.00",
        section="194I",
        payee="Next Quarter Landlord",
        pan="ABCPL1234E",
    )
    return books


def test_the_26q_carries_the_quarters_deductions_as_the_return_lists_them() -> None:
    books = _seeded_quarter()
    tds_return = TdsReturnService(books.session).build(
        books.firm.id, return_quarter("2026-27", "Q1")
    )

    assert (tds_return.tan, tds_return.pan) == ("BLRA12345C", "AAAFA1234B")
    principal, fee = tds_return.deductees
    assert (
        principal.serial,
        principal.section,
        principal.deductee_code,
        principal.pan,
        principal.amount_paid,
        principal.tds_amount,
        principal.rate_percent,
        principal.higher_rate_reason,
    ) == (
        1,
        "194Q",
        "01",
        COMPANY_PAN,
        Decimal("100000.00"),
        Decimal("100.00"),
        Decimal("0.1000"),
        "",
    )
    # No PAN: filed as PANNOTAVBL, reason C, at the higher rate.
    assert (fee.pan, fee.deductee_code, fee.higher_rate_reason, fee.rate_percent) == (
        PAN_NOT_AVAILABLE,
        "",
        "C",
        Decimal("20.0000"),
    )
    assert [
        (due.section, due.month, due.tds_amount, due.due_date)
        for due in tds_return.challans
    ] == [
        ("194Q", "April 2026", Decimal("100.00"), date(2026, 5, 7)),
        ("194J", "May 2026", Decimal("6000.00"), date(2026, 6, 7)),
    ]
    left_out = {item.row.section: item.reason for item in tds_return.excluded}
    assert set(left_out) == {"194C", "192"}
    assert "24Q" in left_out["192"]
    assert "reversed" in left_out["194C"]
    assert len(tds_return.problems) == 1  # the deductee with no PAN


def test_the_26q_file_and_its_csv() -> None:
    books = _seeded_quarter()
    tds_return = TdsReturnService(books.session).build(
        books.firm.id, return_quarter("2026-27", "Q1")
    )
    workbook = load_workbook(BytesIO(return_workbook(tds_return)))
    assert workbook.sheetnames == [
        "Deductor",
        "Deductees",
        "Challans due",
        "Not in this return",
    ]
    deductor = {
        row[0]: row[1]
        for row in workbook["Deductor"].iter_rows(max_row=11, values_only=True)
    }
    assert deductor["TAN"] == "BLRA12345C"
    assert deductor["Assessment year"] == "2027-28"
    rows = list(workbook["Deductees"].iter_rows(values_only=True))
    assert rows[0] == DEDUCTEE_HEADINGS
    assert rows[1][2:6] == ("194Q", "01", COMPANY_PAN, "Vendor One")
    assert rows[2][4] == PAN_NOT_AVAILABLE
    assert workbook["Not in this return"].max_row == 3

    lines = deductees_csv(tds_return).splitlines()
    assert lines[0].split(",")[:3] == ["Sr No", "Challan Serial No", "Section"]
    assert lines[1].split(",")[2:7] == [
        "194Q",
        "01",
        COMPANY_PAN,
        "Vendor One",
        "20-04-2026",
    ]
    assert len(lines) == 3


def _scope(firm_id: UUID) -> ResolvedFirmScope:
    user = uuid4()
    return ResolvedFirmScope(
        principal=Principal(
            subject=user,
            roles=frozenset(),
            permissions=frozenset({"ACCOUNT_VIEW"}),
            claims=TokenClaims(
                sub=str(user),
                type=TokenType.ACCESS,
                iat=1,
                exp=4_102_444_800,
                permissions=["ACCOUNT_VIEW"],
            ),
        ),
        firm_id=firm_id,
    )


def test_the_26q_routes_answer_one_quarter_and_refuse_anything_else() -> None:
    books = _seeded_quarter()
    scope = _scope(books.firm.id)
    page = tds_26q_deductees(
        scope=scope,
        financial_year="2026-27",
        quarter="Q1",
        page=1,
        page_size=1,
        db=books.session,
    )
    assert [row.serial for row in page.data] == [1]
    assert page.pagination.total_records == 2

    response = tds_26q_file(
        scope=scope,
        financial_year="2026-27",
        quarter="Q1",
        format="csv",
        db=books.session,
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="26Q-2026-27-Q1-deductees.csv"'
    )
    with pytest.raises(ValidationError, match="one quarter"):
        tds_26q_file(
            scope=scope,
            financial_year="2026-27",
            quarter="H1",
            format="xlsx",
            db=books.session,
        )
