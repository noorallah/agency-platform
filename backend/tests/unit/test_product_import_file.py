"""Products come in from a file, with a template to fill in (backlog 46).

Every row is checked before anything is written and every problem comes back
with its row and column; references are matched by code or name and reported,
never guessed; the file commits whole or not at all; and an existing product
can be updated by code, a blank cell leaving its field alone.
"""

from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.business.models import BusinessProfile
from app.common.audit.models import AuditLog
from app.core.database.base import Base
from app.firms.models import Firm
from app.products.models import Product, ProductCategory
from app.products.services import ProductService
from app.products.services.brands import BrandService
from app.products.services.product_import import (
    COLUMNS,
    ImportReport,
    template_csv,
    template_workbook,
)
from app.tax.models import TaxProfile, TaxSystem
from app.uom.models import Uom
from tests.unit.test_product_master import TEST_BRAND


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database holding the whole schema."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _firm(session: Session) -> Firm:
    """Add a firm with a category tree, two units and one tax group."""
    firm = Firm(
        name="Import Firm",
        code="IMPORT",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add_all(
        [
            firm,
            BusinessProfile(
                code="GENERIC",
                name="Generic",
                industry_type="GENERIC",
                status="ACTIVE",
                is_default=True,
                default_settings={},
            ),
            Uom(code="PCS", name="Pieces", symbol="pc"),
            Uom(code="KG", name="Kilogram", symbol="kg", dimension="WEIGHT"),
        ]
    )
    session.flush()
    grocery = ProductCategory(
        firm_id=firm.id, code="GROCERY", name="Grocery", level=0, path="GROCERY"
    )
    session.add(grocery)
    session.flush()
    system = TaxSystem(firm_id=firm.id, code="GST", name="GST", display_name="GST")
    session.add_all(
        [
            ProductCategory(
                firm_id=firm.id,
                code="RICE",
                name="Rice and grains",
                parent_id=grocery.id,
                level=1,
                path="GROCERY/RICE",
            ),
            system,
        ]
    )
    session.flush()
    session.add(
        TaxProfile(
            firm_id=firm.id,
            tax_system_id=system.id,
            code="GST5",
            name="GST 5",
            label="GST 5",
            status="ACTIVE",
            group_code="GST5",
        )
    )
    for brand in (TEST_BRAND, "Acme"):
        BrandService(session).ensure_brand(brand, firm_id=firm.id, actor_id=firm.id)
    session.commit()
    return firm


def _csv(*lines: str, brand: str | None = TEST_BRAND) -> bytes:
    """Join lines into a CSV file's bytes.

    A new product needs a brand, so a file whose heading names none gets a
    ``Brand`` column holding ``brand`` on every row; ``brand=None`` leaves
    the file exactly as written.
    """
    if brand is not None and "Brand" not in lines[0].split(","):
        lines = (
            f"{lines[0]},Brand",
            *(f"{line},{brand}" for line in lines[1:]),
        )
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
    file_format: str = "csv",
) -> ImportReport:
    """Check or apply one file as a single actor."""
    return ProductService(session).check_product_file(
        content,
        file_format,  # type: ignore[arg-type]
        firm_scope=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(fresh.scalars(select(Product.code)).all())


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, so the file is fixed once, not row by row."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    ProductService(session).import_products_csv(
        f"Code,Name,Brand\nTAKEN,Already here,{TEST_BRAND}\n",
        firm_scope=firm.id,
        actor_id=uuid4(),
    )
    content = _csv(
        "Code,Name,Category,Unit,TaxGroup,SellingPrice,TrackBatch",
        "GOOD-1,Good,Grocery,PCS,GST5,10,No",
        "BAD-UNIT,Bad unit,GROCERY,BOXES,,,",
        "BAD-CAT,Bad category,Hardware,,,,",
        "BAD-NUM,Bad number,,,,ten,",
        "BAD-FLAG,Bad flag,,,,,maybe",
        "BAD-TAX,Bad tax,,,GST99,,",
        "TAKEN,Clash,,,,,",
        "GOOD-1,Twice,,,,,",
        ",No code,,,,,",
        "NO-NAME,,,,,,",
    )

    report = _run(session, firm.id, content, apply=True)

    found = {(issue.row, issue.column) for issue in report.issues}
    assert found == {
        (3, "Unit"),
        (4, "Category"),
        (5, "SellingPrice"),
        (6, "TrackBatch"),
        (7, "TaxGroup"),
        (8, "Code"),
        (9, "Code"),
        (10, "Code"),
        (11, "Name"),
    }
    assert "row 2 already has it" in next(
        issue.message for issue in report.issues if issue.row == 9
    )
    assert not report.imported
    assert _codes(factory) == ["TAKEN"]


def test_a_clean_file_imports_whole_with_its_references_resolved() -> None:
    """Category, sub category, unit and tax group are matched by code or name."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Code,Name,Category,SubCategory,Unit,HSN,TaxGroup,PurchasePrice,"
        "SellingPrice,MRP,TrackBatch,TrackExpiry",
        "rice-5,Basmati 5 kg,grocery,Rice and grains,kg,10063020,gst5,400,"
        '"1,450",1499,Yes,y',
        "salt-1,Salt,GROCERY,,Pieces,,,,,,,",
    )

    checked = _run(session, firm.id, content, apply=False)
    assert checked.issues == []
    assert (checked.to_create, checked.imported) == (2, False)
    assert _codes(factory) == []

    report = _run(session, firm.id, content, apply=True)

    assert report.imported and report.to_create == 2
    with factory() as fresh:
        rice = fresh.scalars(select(Product).where(Product.code == "RICE-5")).one()
        kg = fresh.scalars(select(Uom).where(Uom.code == "KG")).one()
        sub = fresh.scalars(
            select(ProductCategory).where(ProductCategory.code == "RICE")
        ).one()
        assert rice.sub_category_id == sub.id
        assert rice.category_id == sub.parent_id
        assert (rice.base_uom_id, rice.sales_uom_id, rice.unit) == (
            kg.id,
            kg.id,
            "KG",
        )
        assert rice.tax_profile_group_code == "GST5"
        assert str(rice.selling_price) == "1450.00" or rice.selling_price == 1450
        assert rice.track_batch and rice.track_expiry
        salt = fresh.scalars(select(Product).where(Product.code == "SALT-1")).one()
        assert salt.unit == "PCS"
        audits = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "product.created")
        ).all()
        assert len(audits) == 2


def test_headings_from_other_software_are_read() -> None:
    """Case, spaces, punctuation and common names do not need editing away."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Item Code,ITEM NAME,HSN/SAC,Sale Price,UOM,Opening Qty,Make",
        f"X-1,Imported,1234,25,pcs,10,{TEST_BRAND}",
        brand=None,
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.imported
    assert report.columns_used == [
        "Code",
        "Name",
        "Unit",
        "HSN",
        "Brand",
        "SellingPrice",
    ]
    assert report.columns_ignored == ["Opening Qty"]
    with factory() as fresh:
        row = fresh.scalars(select(Product)).one()
        assert (row.name, row.hsn_sac, row.unit) == ("Imported", "1234", "PCS")


def test_a_file_without_the_required_headings_is_refused_at_row_one() -> None:
    """A sheet with no Code column says so rather than skipping every row."""
    session = _factory()()
    firm = _firm(session)

    report = _run(session, firm.id, _csv("Item,Price", "A,1"), apply=True)

    assert [issue.row for issue in report.issues] == [1]
    assert "Code or Name" in report.issues[0].message


def test_update_by_code_changes_what_the_file_says_and_leaves_blanks() -> None:
    """A re-run corrects a migration; a blank cell is not an instruction."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _run(
        session,
        firm.id,
        _csv("Code,Name,Brand,SellingPrice", "KEEP,Kept,Acme,10"),
        apply=True,
    )
    fix = _csv("Code,Name,Brand,SellingPrice", "KEEP,,,12", "NEW-1,New one,Acme,")

    refused = _run(session, firm.id, fix, apply=True)
    assert [(issue.row, issue.column) for issue in refused.issues] == [(2, "Code")]

    report = _run(session, firm.id, fix, apply=True, existing="update")

    assert report.imported
    assert (report.to_create, report.to_update) == (1, 1)
    with factory() as fresh:
        kept = fresh.scalars(select(Product).where(Product.code == "KEEP")).one()
        assert (kept.name, kept.brand) == ("Kept", "Acme")
        assert kept.selling_price == 12
        updated = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "product.updated")
        ).all()
        assert len(updated) == 1
    assert _codes(factory) == ["KEEP", "NEW-1"]


def test_the_template_round_trips_and_lists_the_firms_codes() -> None:
    """The template's own example row imports, and its lists are this firm's."""
    factory = _factory()
    session = factory()
    firm = _firm(session)

    content = template_workbook(session, firm.id)

    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Products", "Notes", "Lists"]
    headings = [cell.value for cell in workbook["Products"][1]]
    assert headings == [column.heading for column in COLUMNS]
    notes = {row[0] for row in workbook["Notes"].iter_rows(values_only=True)}
    assert {column.heading for column in COLUMNS} <= notes
    listed = {
        value
        for row in workbook["Lists"].iter_rows(values_only=True)
        for value in row
        if value
    }
    assert {"GROCERY", "RICE", "PCS", "KG", "GST5"} <= listed
    assert template_csv().splitlines()[0].split(",")[0] == "Code"

    # The example names GROCERY, PCS and GST5, all of which exist here.
    report = _run(session, firm.id, content, apply=True, file_format="xlsx")
    assert report.issues == [], [issue.describe() for issue in report.issues]
    assert _codes(factory) == ["RICE-5KG"]


def test_a_workbook_with_whole_numbers_reads_them_as_text() -> None:
    """A code typed as 1001 in Excel is 1001, not 1001.0."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Code", "Name", "HSN", "MRP", "Brand"])
    sheet.append([1001, "Numbered", 10063020.0, 99.5, TEST_BRAND])
    buffer = BytesIO()
    workbook.save(buffer)

    report = _run(session, firm.id, buffer.getvalue(), apply=True, file_format="xlsx")

    assert report.imported, [issue.describe() for issue in report.issues]
    with factory() as fresh:
        row = fresh.scalars(select(Product)).one()
        assert (row.code, row.hsn_sac) == ("1001", "10063020")


def _typed_firm(session: Session) -> tuple[Firm, dict[str, UUID]]:
    """Return a firm trading in Medicine and Paint, each with a category.

    TABLETS carries Medicine and STRIPS sits under it with no type of its
    own; ENAMELS carries Paint. Two unit sets: boxes of ten marked for
    Medicine, and loose pieces marked for nothing.
    """
    from app.products.goods_type_seed import seed_goods_types
    from app.products.models.goods_type import FirmGoodsType, GoodsType
    from app.uom.models import UnitSet, UnitSetGoodsType

    firm = _firm(session)
    seed_goods_types(session)
    types = {
        row.code: row.id
        for row in session.scalars(
            select(GoodsType).where(GoodsType.code.in_(("MEDICINE", "PAINT")))
        )
    }
    session.add_all(
        FirmGoodsType(firm_id=firm.id, goods_type_id=type_id)
        for type_id in types.values()
    )
    tablets = ProductCategory(
        firm_id=firm.id,
        code="TABLETS",
        name="Tablets",
        level=0,
        path="TABLETS",
        goods_type_id=types["MEDICINE"],
    )
    session.add_all(
        [
            tablets,
            ProductCategory(
                firm_id=firm.id,
                code="ENAMELS",
                name="Enamels",
                level=0,
                path="ENAMELS",
                goods_type_id=types["PAINT"],
            ),
            Uom(code="BOX", name="Box", symbol="bx"),
        ]
    )
    session.flush()
    session.add(
        ProductCategory(
            firm_id=firm.id,
            code="STRIPS",
            name="Strips",
            parent_id=tablets.id,
            level=1,
            path="TABLETS/STRIPS",
        )
    )
    units = {row.code: row.id for row in session.scalars(select(Uom))}
    boxed = UnitSet(
        name="Piece, box of 10",
        base_uom_id=units["PCS"],
        inventory_uom_id=units["PCS"],
        purchase_uom_id=units["BOX"],
        sales_uom_id=units["PCS"],
        conversion_factor=10,
    )
    loose = UnitSet(name="Piece, loose", base_uom_id=units["PCS"])
    session.add_all([boxed, loose])
    session.flush()
    session.add(UnitSetGoodsType(unit_set_id=boxed.id, goods_type_id=types["MEDICINE"]))
    session.commit()
    return firm, units


def test_a_file_giving_only_the_category_takes_the_goods_types_switches() -> None:
    """The switch columns are optional: the category's goods type fills them."""
    factory = _factory()
    session = factory()
    firm, _units = _typed_firm(session)
    content = _csv(
        "Code,Name,Category,SubCategory,Unit,TrackBatch,TrackExpiry",
        "tab-1,Tablet one,TABLETS,,PCS,,",
        "tab-2,Tablet two,TABLETS,Strips,PCS,,",
        "tab-3,Tablet three,TABLETS,,PCS,Yes,No",
        "enm-1,Enamel one,ENAMELS,,PCS,,",
        "gro-1,Rice,GROCERY,,PCS,,",
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.issues == [] and report.imported
    with factory() as fresh:
        rows = {row.code: row for row in fresh.scalars(select(Product))}
    switches = {
        code: (row.track_batch, row.track_expiry, row.require_batch_on_receipt)
        for code, row in rows.items()
    }
    assert switches == {
        "TAB-1": (True, True, True),
        # A sub category with no type of its own takes its parent's.
        "TAB-2": (True, True, True),
        # A switch the file names is the file's, on or off.
        "TAB-3": (True, False, True),
        "ENM-1": (True, False, True),
        # No goods type is General: nothing tracked.
        "GRO-1": (False, False, False),
    }
    assert rows["TAB-1"].goods_type_id == rows["TAB-2"].goods_type_id
    assert rows["GRO-1"].goods_type_id is None


def test_a_unit_set_column_fills_the_units_and_warns_of_another_types_set() -> None:
    """One column for the units and the pack; a mismatch is said, not refused."""
    from app.uom.models import ConversionRule

    factory = _factory()
    session = factory()
    firm, units = _typed_firm(session)
    content = _csv(
        "Code,Name,Category,Unit,UnitSet",
        'tab-1,Tablet one,TABLETS,,"Piece, box of 10"',
        'enm-1,Enamel one,ENAMELS,,"piece, box of 10"',
        'enm-2,Enamel two,ENAMELS,,"Piece, loose"',
        'enm-3,Enamel three,ENAMELS,KG,"Piece, box of 10"',
    )

    checked = _run(session, firm.id, content, apply=False)

    # Nothing is refused: the check stays clean and says the one odd pairing
    # once per row it is on.
    assert checked.issues == []
    assert [(item.row, item.code, item.column) for item in checked.warnings] == [
        (3, "ENM-1", "UnitSet"),
        (5, "ENM-3", "Unit"),
        (5, "ENM-3", "UnitSet"),
    ]
    assert "other goods types" in checked.warnings[0].message
    # A row naming a Unit beside the set is told its Unit was passed over.
    assert "is passed over" in checked.warnings[1].message
    assert "Piece, box of 10" in checked.warnings[1].message
    assert _codes(factory) == []

    report = _run(session, firm.id, content, apply=True)

    assert report.imported and report.to_create == 4
    with factory() as fresh:
        rows = {row.code: row for row in fresh.scalars(select(Product))}
        rules = {
            rule.product_id: rule
            for rule in fresh.scalars(
                select(ConversionRule).where(ConversionRule.is_deleted.is_(False))
            )
            if rule.product_id is not None
        }
    # The set speaks for every unit, a Unit on the row or not: kept, the
    # Unit made the purchase unit the stock unit and the pack rule was
    # dropped without a word (D-MST-17).
    for code in ("TAB-1", "ENM-1", "ENM-3"):
        assert (rows[code].base_uom_id, rows[code].purchase_uom_id) == (
            units["PCS"],
            units["BOX"],
        )
        assert rules[rows[code].id].conversion_factor == 10
    assert rows["ENM-2"].purchase_uom_id is None and rows["ENM-2"].id not in rules


def test_an_unknown_unit_set_is_a_problem_and_an_existing_product_keeps_its_units() -> (
    None
):
    """A name that is no set stops the file; a set on an update is passed over."""
    factory = _factory()
    session = factory()
    firm, units = _typed_firm(session)
    _run(
        session,
        firm.id,
        _csv("Code,Name,Category,Unit", "tab-1,Tablet one,TABLETS,PCS"),
        apply=True,
    )

    unknown = _run(
        session,
        firm.id,
        _csv("Code,Name,UnitSet", "tab-9,Tablet nine,Barrel of 40"),
        apply=True,
    )
    assert [(item.column, item.message) for item in unknown.issues] == [
        ("UnitSet", "'Barrel of 40' is not an active unit set.")
    ]
    assert not unknown.imported

    updated = _run(
        session,
        firm.id,
        _csv("Code,Name,UnitSet", 'tab-1,Tablet one renamed,"Piece, box of 10"'),
        apply=True,
        existing="update",
    )
    assert updated.issues == [] and updated.imported
    assert [item.column for item in updated.warnings] == ["UnitSet"]
    assert "passed over" in updated.warnings[0].message
    with factory() as fresh:
        row = fresh.scalars(select(Product).where(Product.code == "TAB-1")).one()
    assert (row.name, row.purchase_uom_id) == ("Tablet one renamed", units["PCS"])


def test_a_new_product_row_with_a_blank_brand_is_reported_on_the_column() -> None:
    """A product always has a brand: a blank cell on a new row is a problem."""
    factory = _factory()
    session = factory()
    firm = _firm(session)

    report = _run(
        session,
        firm.id,
        _csv("Code,Name,Brand", "NEW-1,No brand,", f"NEW-2,Branded,{TEST_BRAND}"),
        apply=True,
    )

    assert [(item.row, item.column, item.message) for item in report.issues] == [
        (2, "Brand", "is required.")
    ]
    assert not report.imported
    assert _codes(factory) == []


def test_an_unknown_brand_is_reported_on_the_column() -> None:
    """The file never makes a brand: it must sit under a company already."""
    session = _factory()()
    firm = _firm(session)

    report = _run(
        session,
        firm.id,
        _csv("Code,Name,Brand", "NEW-1,Strange,Zenith"),
        apply=True,
    )

    assert [(item.row, item.column, item.message) for item in report.issues] == [
        (
            2,
            "Brand",
            "'Zenith' is not one of the firm's brands; add it under its "
            "company first.",
        )
    ]
    assert not report.imported


def test_a_blank_brand_on_an_existing_product_leaves_the_stored_brand() -> None:
    """A re-run that only fixes a price does not need the brand repeated."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _run(
        session,
        firm.id,
        _csv("Code,Name,Brand", "KEEP,Kept,Acme"),
        apply=True,
    )

    report = _run(
        session,
        firm.id,
        _csv("Code,Name,Brand", "KEEP,Kept again,"),
        apply=True,
        existing="update",
    )

    assert report.issues == [] and report.imported
    with factory() as fresh:
        kept = fresh.scalars(select(Product).where(Product.code == "KEEP")).one()
        assert (kept.name, kept.brand) == ("Kept again", "Acme")
        assert kept.brand_id is not None


def test_the_templates_lists_sheet_names_the_firms_brands() -> None:
    """The Brand column of the Lists sheet is the firm's own brands."""
    session = _factory()()
    firm = _firm(session)

    workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))

    rows = list(workbook["Lists"].iter_rows(values_only=True))
    column = next(index for index, value in enumerate(rows[0]) if value == "Brand")
    assert [row[column] for row in rows[1:] if row[column]] == ["Acme", TEST_BRAND]


def test_pieces_per_pack_makes_a_new_products_box() -> None:
    """The old program's pieces per case becomes the item's pack (L5, G4).

    Above one, a new product is kept in its Unit (PCS where blank) and bought
    and sold in boxes of that many; one is pieces only; a product that exists
    and a row with a unit set are said to be passed over, never refused.
    """
    from app.uom.services.packs import pack_of

    factory = _factory()
    session = factory()
    firm, units = _typed_firm(session)
    _run(
        session,
        firm.id,
        _csv("Code,Name,Category,Unit", "old-1,Old one,TABLETS,PCS"),
        apply=True,
    )
    content = _csv(
        "Code,Name,Category,Unit,PiecesPerCase,PackUnit",
        "pk-1,Pack one,TABLETS,,12,",
        "pk-2,Pack two,TABLETS,KG,24,box",
        "pk-3,Pack three,TABLETS,PCS,1,",
        "pk-4,Pack four,TABLETS,PCS,,",
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.issues == [] and report.warnings == [] and report.imported
    with factory() as fresh:
        rows = {row.code: row for row in fresh.scalars(select(Product))}
        packs = {
            code: pack_of(fresh, rows[code], side=side, on_date=date(2030, 1, 1))
            for code in ("PK-1", "PK-2", "PK-3", "PK-4")
            for side in ("SELL",)
        }
        buying = pack_of(fresh, rows["PK-1"], side="BUY", on_date=date(2030, 1, 1))
    assert (rows["PK-1"].base_uom_id, rows["PK-1"].purchase_uom_id) == (
        units["PCS"],
        units["BOX"],
    )
    assert rows["PK-1"].sales_uom_id == units["PCS"]
    assert packs["PK-1"] is not None and packs["PK-1"].size == 12
    assert buying is not None and (buying.size, buying.unit_code) == (12, "BOX")
    assert rows["PK-2"].base_uom_id == units["KG"]
    assert packs["PK-2"] is not None and packs["PK-2"].size == 24
    assert packs["PK-3"] is None and rows["PK-3"].purchase_uom_id == units["PCS"]
    assert packs["PK-4"] is None

    passed = _run(
        session,
        firm.id,
        _csv(
            "Code,Name,Category,Unit,UnitSet,PiecesPerPack,PackUnit",
            "old-1,Old one,TABLETS,PCS,,6,",
            'pk-5,Pack five,TABLETS,,"Piece, box of 10",6,',
            "pk-6,Pack six,TABLETS,PCS,,,BOX",
        ),
        apply=True,
        existing="update",
    )
    assert passed.issues == [] and passed.imported
    assert [(item.code, item.column) for item in passed.warnings] == [
        ("OLD-1", "PiecesPerPack"),
        ("PK-5", "PiecesPerPack"),
        ("PK-6", "PackUnit"),
    ]
    assert all("is passed over" in item.message for item in passed.warnings)
    with factory() as fresh:
        old = fresh.scalars(select(Product).where(Product.code == "OLD-1")).one()
        five = fresh.scalars(select(Product).where(Product.code == "PK-5")).one()
        assert pack_of(fresh, old, side="SELL", on_date=date(2030, 1, 1)) is None
        five_pack = pack_of(fresh, five, side="SELL", on_date=date(2030, 1, 1))
    # The set's factor stands, not the row's six.
    assert five_pack is not None and five_pack.size == 10


def test_a_wrong_pieces_per_pack_is_a_problem_on_its_column() -> None:
    """A size that is no whole number, an unknown box, or the piece itself."""
    factory = _factory()
    session = factory()
    firm, _units = _typed_firm(session)

    report = _run(
        session,
        firm.id,
        _csv(
            "Code,Name,Category,Unit,PiecesPerPack,PackUnit",
            "bad-1,Bad one,TABLETS,PCS,2.5,",
            "bad-2,Bad two,TABLETS,PCS,0,",
            "bad-3,Bad three,TABLETS,PCS,many,",
            "bad-4,Bad four,TABLETS,PCS,12,CRATE",
            "bad-5,Bad five,TABLETS,PCS,12,PCS",
        ),
        apply=True,
    )

    assert not report.imported and _codes(factory) == []
    assert [(item.code, item.column) for item in report.issues] == [
        ("BAD-1", "PiecesPerPack"),
        ("BAD-2", "PiecesPerPack"),
        ("BAD-3", "PiecesPerPack"),
        ("BAD-4", "PackUnit"),
        ("BAD-5", "PackUnit"),
    ]
    assert "'CRATE' is not a unit" in report.issues[3].message
