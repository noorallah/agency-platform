"""Ledger accounts come in from a file, with a template to fill in (Swathi G6).

The file rules are the other masters', from ``app.common.file_import``. What
these tests pin is what is particular to the chart of accounts -- the code is
the key, every account sits under a group the firm already has (named by code
or name) and takes its type, and every row goes through the finance service's
own guards before the file commits once.

The sessions do not autoflush, because a request's do not.
"""

import sqlite3
from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import Connection, create_engine, event, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.finance.models.finance import AccountGroup, LedgerAccount
from app.finance.services.finance_service import FinanceService
from app.finance.services.ledger_account_import import (
    COLUMNS,
    LedgerAccountFileImporter,
    template_csv,
    template_workbook,
)
from app.firms.models import Firm
from app.imports.services import IMPORT_KINDS, columns_for_kind


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush.

    The finance service inserts each account inside a savepoint. The sqlite3
    driver opens no transaction before a ``SAVEPOINT``, so releasing the
    first one would commit it -- which PostgreSQL never does, and which would
    make a rolled-back check look as if it wrote. SQLAlchemy's documented
    recipe makes the driver begin the transaction itself, as PostgreSQL's
    does.
    """
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _no_driver_transactions(connection: sqlite3.Connection, _: object) -> None:
        """Stop the driver opening transactions on its own."""
        connection.isolation_level = None

    @event.listens_for(engine, "begin")
    def _begin(connection: Connection) -> None:
        """Begin each transaction explicitly, so a savepoint nests inside it."""
        connection.exec_driver_sql("BEGIN")

    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with four account groups and one account already charted."""
    firm = Firm(
        name="Chart Firm",
        code="CHART",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    groups = {
        code: AccountGroup(
            firm_id=firm.id, code=code, name=name, account_type=account_type
        )
        for code, name, account_type in (
            ("SALES", "Sales Accounts", "INCOME"),
            ("INDEXP", "Indirect Expenses", "EXPENSE"),
            ("ADMEXP", "Indirect Expenses", "EXPENSE"),
            ("CASH", "Cash-in-hand", "ASSET"),
        )
    }
    session.add_all(groups.values())
    session.flush()
    session.add(
        LedgerAccount(
            firm_id=firm.id,
            account_group_id=groups["CASH"].id,
            code="1000",
            name="Cash",
            account_type="ASSET",
            is_balance_sheet=True,
            is_profit_loss=False,
        )
    )
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
) -> ImportReport[LedgerAccount]:
    """Check or apply one file as a single actor."""
    return LedgerAccountFileImporter(session, FinanceService(session)).run(
        content,
        file_format="csv",
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(
            fresh.scalars(
                select(LedgerAccount.code).where(LedgerAccount.is_deleted.is_(False))
            ).all()
        )


def _account(factory: sessionmaker[Session], code: str) -> LedgerAccount:
    """Read one stored account by code, from a session of its own."""
    with factory() as fresh:
        return fresh.scalars(
            select(LedgerAccount).where(LedgerAccount.code == code)
        ).one()


def test_the_template_names_the_columns_and_lists_groups_and_accounts() -> None:
    """The CSV carries the headings; the workbook lists the groups to name."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        assert template_csv().splitlines()[0] == ",".join(c.heading for c in COLUMNS)
        workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))
        assert workbook.sheetnames == ["Ledger accounts", "Notes", "Lists"]
        lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
        assert lists[0] == [
            "Group code",
            "Group name",
            "Group type",
            "Account code",
            "Account name",
        ]
        assert {row[0] for row in lists[1:]} == {"SALES", "INDEXP", "ADMEXP", "CASH"}
        assert lists[1][3:] == ["1000", "Cash"]


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        report = _run(
            session,
            firm.id,
            _csv(
                "Ledger Code,Ledger Name,Under,Cost Centre,Active",
                "4100,Sales - Local,SALES,No,Yes",
                "4200,,SALES,,",
                "5100,Rent,,,",
                "5200,Postage,Postage Group,,",
                "5300,Travel,Indirect Expenses,,",
                "5400,Phone,indexp,maybe,",
                "4100,Again,SALES,,",
                ",Nameless,SALES,,",
                "1000,Cash Again,CASH,,",
                "BAD CODE,Spaced,SALES,,",
                "5500,Service Income,sales accounts,,",
            ),
            apply=False,
        )
    found = {(issue.row, issue.column, issue.message) for issue in report.issues}
    assert (3, "Name", "is needed for a new ledger account.") in found
    assert (4, "Group", "is needed for a new ledger account.") in found
    assert any(
        row == 5 and column == "Group" and "not one of the firm's" in message
        for row, column, message in found
    )
    # Two groups share the name, so the name alone cannot say which.
    assert any(
        row == 6 and column == "Group" and "ADMEXP, INDEXP" in message
        for row, column, message in found
    )
    assert (7, "Cost Centre", "'maybe' should be Yes or No.") in found
    assert (8, "Code", "appears again; row 2 already has it.") in found
    assert (9, "Code", "is required.") in found
    assert any(row == 10 and column == "Code" for row, column, _ in found)
    assert any(row == 11 and column == "Code" for row, column, _ in found)
    assert len(report.issues) == 9
    # 4100 and 5500 (by its group's name, whatever the case) staged, then
    # rolled back with the rest.
    assert report.to_create == 2
    assert report.imported is False
    assert _codes(factory) == ["1000"]


def test_an_apply_is_all_or_nothing() -> None:
    """One bad row keeps the good ones out; a clean file goes in whole."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        bad = _run(
            session,
            firm.id,
            _csv("Code,Name,Group", "4100,Sales - Local,SALES", "5100,Rent,NOPE"),
            apply=True,
        )
        assert bad.imported is False
        assert _codes(factory) == ["1000"]

        good = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Group,Description,Cost Centre,Profit Centre,Active",
                "4100,Sales - Local,sales accounts,Counter sales,,,",
                "5100,Rent,INDEXP,,Yes,No,No",
            ),
            apply=True,
        )
    assert good.imported is True
    assert good.to_create == 2
    assert _codes(factory) == ["1000", "4100", "5100"]
    sales = _account(factory, "4100")
    rent = _account(factory, "5100")
    # The account takes its group's type, and the statement follows the type.
    assert sales.account_type == "INCOME"
    assert sales.is_profit_loss is True
    assert sales.is_balance_sheet is False
    assert sales.description == "Counter sales"
    assert sales.is_active is True
    assert rent.account_type == "EXPENSE"
    assert rent.requires_cost_center is True
    assert rent.requires_profit_center is False
    assert rent.is_active is False
    with factory() as fresh:
        # Each account is audited as the form audits it.
        created = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "finance.ledger_account.created")
        ).all()
        assert len(created) == 2


def test_a_second_apply_refuses_the_same_codes_unless_asked_to_update() -> None:
    """A code already an account is refused; update changes only what is sent."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        first = _run(
            session,
            firm.id,
            _csv("Code,Name,Group,Description", "5100,Rent,INDEXP,Shop rent"),
            apply=True,
        )
        assert first.imported is True

        again = _run(session, firm.id, _csv("Code,Name", "5100,Rent Paid"), apply=True)
        assert again.imported is False
        assert [(i.row, i.column) for i in again.issues] == [(2, "Code")]
        assert "is already a ledger account" in again.issues[0].message

        # An account cannot move to a group of another type.
        moved = _run(
            session,
            firm.id,
            _csv("Code,Group", "5100,SALES"),
            apply=True,
            existing="update",
        )
        assert moved.imported is False
        assert [(i.row, i.column) for i in moved.issues] == [(2, "Group")]

        updated = _run(
            session,
            firm.id,
            _csv("Code,Name,Group,Active", "5100,,ADMEXP,No"),
            apply=True,
            existing="update",
        )
    assert updated.imported is True
    assert updated.to_update == 1
    rent = _account(factory, "5100")
    # A blank cell or a missing column keeps what the account held.
    assert rent.name == "Rent"
    assert rent.description == "Shop rent"
    assert rent.is_active is False
    with factory() as fresh:
        admexp = fresh.scalars(
            select(AccountGroup).where(AccountGroup.code == "ADMEXP")
        ).one()
        assert rent.account_group_id == admexp.id
        updates = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "finance.ledger_account.updated")
        ).all()
        assert len(updates) == 1


def test_the_chart_can_be_mapped_from_another_softwares_file() -> None:
    """The kind is registered under the chart's own code, with its columns."""
    factory = _factory()
    assert IMPORT_KINDS["ledger-accounts"].permission == "ACCOUNT_MANAGE"
    with factory() as session:
        assert columns_for_kind(session, "ledger-accounts") == COLUMNS
