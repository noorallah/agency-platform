"""Salesmen come in from a file, with a template to fill in (Swathi G3).

The file rules are the other masters', from ``app.common.file_import``. What
these tests pin is what is particular to salesmen -- the code is the key, a
name is unique on the list whatever its case, and every row goes through the
service's own guards before the file commits once.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.firms.models import Firm
from app.imports.services import IMPORT_KINDS, columns_for_kind
from app.salesmen.models import Salesman
from app.salesmen.schemas import SalesmanCreate, SalesmanUpdate
from app.salesmen.services import SalesmanService
from app.salesmen.services.salesman_import import (
    COLUMNS,
    SalesmanFileImporter,
    template_csv,
    template_workbook,
)


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with one salesman already on its list."""
    firm = Firm(
        name="Salesman Firm",
        code="SALESM",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    session.add(Salesman(firm_id=firm.id, code="S00", name="Office"))
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
) -> ImportReport[Salesman]:
    """Check or apply one file as a single actor."""
    return SalesmanFileImporter(session, SalesmanService(session)).run(
        content,
        file_format="csv",
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str | None]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(
            fresh.scalars(
                select(Salesman.code).where(Salesman.is_deleted.is_(False))
            ).all(),
            key=str,
        )


def test_the_template_names_the_columns_and_lists_the_salesmen() -> None:
    """The CSV carries the headings; the workbook lists who is on the list."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        assert template_csv().splitlines()[0] == ",".join(c.heading for c in COLUMNS)
        workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))
        assert workbook.sheetnames == ["Salesmen", "Notes", "Lists"]
        lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
        assert lists[0] == ["Salesman code", "Salesman name"]
        assert lists[1] == ["S00", "Office"]


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        report = _run(
            session,
            firm.id,
            _csv(
                "Salesman Code,Salesman Name,Mobile,Commission,Active",
                "S01,Ravi Kumar,9876543210,Yes,Yes",
                "S02,,,,",
                "S03,Anil,,maybe,",
                "s01,Again,,,",
                ",Nameless,,,",
                "S04,ravi kumar,,,",
                "S05,Phone Order,,No,",
                "S00,Office Again,,,",
            ),
            apply=False,
        )
    found = {(issue.row, issue.column, issue.message) for issue in report.issues}
    assert (3, "Name", "is needed for a new salesman.") in found
    assert (4, "Commission", "'maybe' should be Yes or No.") in found
    assert (5, "Code", "appears again; row 2 already has it.") in found
    assert (6, "Code", "is required.") in found
    # A name already on the list is refused whatever its case.
    assert any(row == 7 for row, _, _ in found)
    assert any(row == 9 and column == "Code" for row, column, _ in found)
    assert len(report.issues) == 6
    assert report.to_create == 2  # S01 and S05 staged, then rolled back
    assert report.imported is False
    assert _codes(factory) == ["S00"]


def test_an_apply_is_all_or_nothing() -> None:
    """One bad row keeps the good ones out; a clean file goes in whole."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        bad = _run(
            session,
            firm.id,
            _csv("Code,Name,Active", "S01,Ravi Kumar,", "S02,Anil,maybe"),
            apply=True,
        )
        assert bad.imported is False
        assert _codes(factory) == ["S00"]

        good = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Phone,Commission,Active,Notes",
                "s01,Ravi Kumar,9876543210,,,North beat",
                "S05,Phone Order,,No,No,",
            ),
            apply=True,
        )
    assert good.imported is True
    assert good.to_create == 2
    assert _codes(factory) == ["S00", "S01", "S05"]
    with factory() as fresh:
        ravi = fresh.scalars(select(Salesman).where(Salesman.code == "S01")).one()
        phone = fresh.scalars(select(Salesman).where(Salesman.code == "S05")).one()
        assert ravi.phone == "9876543210"
        assert ravi.notes == "North beat"
        assert ravi.earns_commission is True
        assert ravi.is_active is True
        assert ravi.user_id is None
        assert phone.earns_commission is False
        assert phone.is_active is False
        # Each salesman is audited as the form audits him.
        assert (
            len(
                fresh.scalars(
                    select(AuditLog).where(AuditLog.action == "salesman.created")
                ).all()
            )
            == 2
        )


def test_a_second_apply_refuses_the_same_codes_unless_asked_to_update() -> None:
    """A code already a salesman is refused; update changes only what is sent."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        first = _run(
            session,
            firm.id,
            _csv("Code,Name,Phone", "S01,Ravi Kumar,9876543210"),
            apply=True,
        )
        assert first.imported is True

        again = _run(session, firm.id, _csv("Code,Name", "s01,Ravi K"), apply=True)
        assert again.imported is False
        assert [(i.row, i.column) for i in again.issues] == [(2, "Code")]
        assert "is already a salesman" in again.issues[0].message

        updated = _run(
            session,
            firm.id,
            _csv("Code,Name,Active", "S01,,No"),
            apply=True,
            existing="update",
        )
    assert updated.imported is True
    assert updated.to_update == 1
    with factory() as fresh:
        ravi = fresh.scalars(select(Salesman).where(Salesman.code == "S01")).one()
        # A blank cell or a missing column keeps what the salesman held.
        assert ravi.name == "Ravi Kumar"
        assert ravi.phone == "9876543210"
        assert ravi.is_active is False
        assert (
            len(
                fresh.scalars(
                    select(AuditLog).where(AuditLog.action == "salesman.updated")
                ).all()
            )
            == 1
        )


def test_create_and_update_still_commit_on_their_own() -> None:
    """The form's path is the stage method and a commit, as before."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        service = SalesmanService(session)
        saved = service.create(
            SalesmanCreate(code="s09", name="Suresh"),
            firm_id=firm.id,
            actor_id=uuid4(),
        )
        service.update(
            saved.id,
            SalesmanUpdate(phone="12345"),
            firm_id=firm.id,
            actor_id=uuid4(),
        )
    assert saved.code == "S09"
    assert _codes(factory) == ["S00", "S09"]
    with factory() as fresh:
        suresh = fresh.scalars(select(Salesman).where(Salesman.code == "S09")).one()
        assert suresh.phone == "12345"


def test_salesmen_can_be_mapped_from_another_softwares_file() -> None:
    """The kind is registered under the salesman code, with its columns."""
    factory = _factory()
    assert IMPORT_KINDS["salesmen"].permission == "SALESMAN_MANAGE"
    with factory() as session:
        assert columns_for_kind(session, "salesmen") == COLUMNS
