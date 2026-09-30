"""Suppliers come in from a file, with a template to fill in (backlog 46).

The file rules are the product and customer imports', from
``app.common.file_import``. What these tests pin is what is particular to
suppliers -- an address placed in the geography masters and never guessed;
one address, contact and bank account per row, updated in place; and a bank
account held to the manage-bank-details duty.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.firms.models import Firm
from app.sales.models.territory import (
    GeoCity,
    GeoCountry,
    GeoDistrict,
    GeoPostalCode,
    GeoState,
)
from app.vendors.models import Vendor, VendorAddress, VendorBankAccount, VendorCategory
from app.vendors.schemas import VendorCreate
from app.vendors.services import VendorService
from app.vendors.services.vendor_import import (
    COLUMNS,
    VendorFileImporter,
    template_csv,
    template_workbook,
)


class _Places:
    """One country, one state, two cities; Chennai has PIN 600001."""

    def __init__(self, session: Session) -> None:
        """Seed the geography masters the addresses are placed in."""
        self.country = GeoCountry(code="IN", name="India", iso2="IN")
        session.add(self.country)
        session.flush()
        self.state = GeoState(country_id=self.country.id, code="TN", name="Tamil Nadu")
        session.add(self.state)
        session.flush()
        self.district = GeoDistrict(state_id=self.state.id, code="CHN", name="Chennai")
        session.add(self.district)
        session.flush()
        self.city = GeoCity(district_id=self.district.id, code="CHN", name="Chennai")
        self.madurai = GeoCity(district_id=self.district.id, code="MDU", name="Madurai")
        session.add_all([self.city, self.madurai])
        session.flush()
        self.pin = GeoPostalCode(city_id=self.city.id, postal_code="600001")
        session.add(self.pin)
        session.commit()


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with one supplier category."""
    firm = Firm(
        name="Import Firm",
        code="IMPORT",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    session.add(VendorCategory(firm_id=firm.id, code="PHARMA", name="Pharma"))
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
    file_format: str = "csv",
    may_manage_bank_details: bool = True,
) -> ImportReport[Vendor]:
    """Check or apply one file as a single actor."""
    return VendorFileImporter(
        session,
        VendorService(session),
        may_manage_bank_details=may_manage_bank_details,
    ).run(
        content,
        file_format=file_format,  # type: ignore[arg-type]
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(fresh.scalars(select(Vendor.code)).all())


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _Places(session)
    content = _csv(
        "Code,Name,Category,Mobile,Email,City,State,PIN,BankName,AccountNumber",
        "GOOD,Good,Pharma,9876543210,a@b.example,,,,,",
        "BAD-CAT,Bad category,Hardware,,,,,,,",
        "BAD-PHONE,Bad phone,,123,,,,,,",
        "BAD-MAIL,Bad email,,,nope,,,,,",
        "BAD-STATE,Bad state,,,,,Atlantis,,,",
        "BAD-PIN,Bad PIN,,,,,,999999,,",
        "NO-LINE,Address without a line,,,,Chennai,,,,",
        "HALF-BANK,Half a bank,,,,,,,SBI,",
        "GOOD,Twice,,,,,,,,",
    )

    report = _run(session, firm.id, content, apply=True)

    found = {(issue.row, issue.column) for issue in report.issues}
    assert found == {
        (3, "Category"),
        (4, "Mobile"),
        (5, "Email"),
        (6, "State"),
        (6, "Address1"),
        (7, "PIN"),
        (7, "Address1"),
        (8, "Address1"),
        (9, "AccountName"),
        (9, "AccountNumber"),
        (10, "Code"),
    }, [issue.describe() for issue in report.issues]
    assert "geography masters" in next(
        issue.message for issue in report.issues if issue.row == 7
    )
    assert _codes(factory) == []


def test_an_address_is_placed_by_its_pin_in_the_geography_masters() -> None:
    """The PIN alone finds the city, district, state and country above it."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    places = _Places(session)
    content = _csv(
        "Supplier Code,Supplier Name,GSTIN/UIN,Address,Pincode,Contact Person,"
        "Bank,Account Holder,Account No,IFSC Code",
        "ACME,Acme Distributors,33aaaca1234a1z5,14 Industrial Estate,600001,"
        "S. Rao,State Bank,Acme Distributors,0012 3456 789,sbin0001234",
        "MDU,Madurai Traders,,1 Temple Road,,,,,,",
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    with factory() as fresh:
        acme = fresh.scalars(select(Vendor).where(Vendor.code == "ACME")).one()
        assert acme.gstin == "33AAACA1234A1Z5"
        assert acme.gst_registration is True
        address = fresh.scalars(
            select(VendorAddress).where(VendorAddress.vendor_id == acme.id)
        ).one()
        assert (
            address.postal_code_id,
            address.city_id,
            address.district_id,
            address.state_id,
            address.country_id,
        ) == (
            places.pin.id,
            places.city.id,
            places.district.id,
            places.state.id,
            places.country.id,
        )
        account = fresh.scalars(select(VendorBankAccount)).one()
        assert (account.account_number, account.ifsc) == (
            "00123456789",
            "SBIN0001234",
        )
        madurai = fresh.scalars(select(Vendor).where(Vendor.code == "MDU")).one()
        plain = fresh.scalars(
            select(VendorAddress).where(VendorAddress.vendor_id == madurai.id)
        ).one()
        assert plain.city_id is None and plain.address_line1 == "1 Temple Road"


def test_a_city_in_one_state_is_placed_by_name() -> None:
    """A city name the masters hold once is enough without a PIN."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    places = _Places(session)
    content = _csv("Code,Name,Address1,City", "MDU,Madurai Traders,1 Road,madurai")

    report = _run(session, firm.id, content, apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    with factory() as fresh:
        address = fresh.scalars(select(VendorAddress)).one()
        assert (address.city_id, address.state_id) == (
            places.madurai.id,
            places.state.id,
        )


def test_a_bank_account_needs_the_bank_details_permission() -> None:
    """Where a supplier is paid is its own duty, from a file as on the form."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Code,Name,BankName,AccountName,AccountNumber",
        "PAYME,Pay Me,SBI,Pay Me,123456",
        "PLAIN,Plain,,,",
    )

    refused = _run(session, firm.id, content, apply=True, may_manage_bank_details=False)

    assert [(issue.row, issue.code) for issue in refused.issues] == [(2, "PAYME")]
    assert "VENDOR_MANAGE_BANK_DETAILS" in refused.issues[0].message
    assert _codes(factory) == []


def test_updating_by_code_changes_only_what_the_file_fills() -> None:
    """A blank cell keeps its field; the primary address changes in place."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    places = _Places(session)
    VendorService(session).create(
        VendorCreate.model_validate(
            {
                "code": "KEEP",
                "name": "Keep Me",
                "email": "keep@supplier.example",
                "addresses": [
                    {
                        "address_type": "BILLING",
                        "address_line1": "1 Old Street",
                        "is_primary": True,
                    },
                    {"address_type": "WAREHOUSE", "address_line1": "Godown 4"},
                ],
            }
        ),
        firm_id=firm.id,
        actor_id=uuid4(),
    )
    content = _csv("Code,Name,Email,Mobile,PIN", "KEEP,,,9876543210,600001")

    refused = _run(session, firm.id, content, apply=True)
    assert "update existing suppliers" in refused.issues[0].message

    report = _run(session, firm.id, content, apply=True, existing="update")

    assert report.imported, [issue.describe() for issue in report.issues]
    assert (report.to_create, report.to_update) == (0, 1)
    with factory() as fresh:
        kept = fresh.scalars(select(Vendor)).one()
        assert (kept.name, kept.email, kept.mobile) == (
            "Keep Me",
            "keep@supplier.example",
            "+919876543210",
        )
        addresses = fresh.scalars(
            select(VendorAddress)
            .where(VendorAddress.is_deleted.is_(False))
            .order_by(VendorAddress.address_line1)
        ).all()
        assert [(a.address_line1, a.postal_code_id) for a in addresses] == [
            ("1 Old Street", places.pin.id),
            ("Godown 4", None),
        ]


def test_the_template_round_trips_and_lists_the_firms_categories() -> None:
    """The template's own example row imports, and its lists are this firm's."""
    factory = _factory()
    session = factory()
    firm = _firm(session)

    content = template_workbook(session, firm.id)

    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Suppliers", "Notes", "Lists"]
    headings = [cell.value for cell in workbook["Suppliers"][1]]
    assert headings == [column.heading for column in COLUMNS]
    listed = {
        value
        for row in workbook["Lists"].iter_rows(values_only=True)
        for value in row
        if value
    }
    assert {"PHARMA", "Pharma", "ACTIVE"} <= listed
    assert template_csv().splitlines()[0].split(",")[0] == "Code"

    report = _run(session, firm.id, content, apply=True, file_format="xlsx")
    assert report.issues == [], [issue.describe() for issue in report.issues]
    assert _codes(factory) == ["SUP-001"]
