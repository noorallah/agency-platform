"""Routes (beats) come in from a file, with a template to fill in (Swathi G3).

The file rules are the other masters', from ``app.common.file_import``. What
these tests pin is what is particular to routes -- a row is a node one level
below the area it names with a round of its own, the area and the route type
must already exist, and every row goes through the territory service's own
guards before the file commits once.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.firms.models import Firm
from app.imports.services import IMPORT_KINDS, columns_for_kind
from app.sales.models import (
    RouteTypeMaster,
    SalesTerritoryNode,
    TerritoryRouteProfile,
    TerritoryWorkingDay,
)
from app.sales.schemas import TerritoryCreate
from app.sales.services import SalesTerritoryService
from app.sales.services.route_import import (
    COLUMNS,
    RouteFileImporter,
    parse_days,
    template_csv,
    template_workbook,
)

_ACTOR = uuid4()


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with Region > Territory > Route, one region, two areas."""
    firm = Firm(
        name="Route Firm",
        code="ROUTES",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.commit()
    service = SalesTerritoryService(session)
    levels = service.get_hierarchy(firm_scope=firm.id, actor_id=_ACTOR).levels
    by_order = {level.level_order: level.id for level in levels}
    region = service.create_territory(
        TerritoryCreate(code="SOUTH", name="South", hierarchy_level_id=by_order[1]),
        firm_scope=firm.id,
        actor_id=_ACTOR,
    )
    for code, name in (("CITY", "City"), ("RURAL", "Rural")):
        service.create_territory(
            TerritoryCreate(
                code=code,
                name=name,
                hierarchy_level_id=by_order[2],
                parent_id=region.id,
            ),
            firm_scope=firm.id,
            actor_id=_ACTOR,
        )
    session.add(RouteTypeMaster(firm_id=firm.id, code="URBAN", name="Urban beat"))
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
) -> ImportReport[SalesTerritoryNode]:
    """Check or apply one file as a single actor."""
    return RouteFileImporter(session, SalesTerritoryService(session)).run(
        content,
        file_format="csv",
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _routes(factory: sessionmaker[Session]) -> list[str]:
    """Read the codes of the stored routes, from a session of its own."""
    with factory() as fresh:
        return sorted(
            fresh.scalars(
                select(SalesTerritoryNode.code)
                .join(
                    TerritoryRouteProfile,
                    TerritoryRouteProfile.territory_id == SalesTerritoryNode.id,
                )
                .where(
                    SalesTerritoryNode.is_deleted.is_(False),
                    TerritoryRouteProfile.is_deleted.is_(False),
                )
            ).all()
        )


def test_the_template_names_the_columns_and_lists_areas_and_types() -> None:
    """The CSV carries the headings; the workbook lists areas and route types."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        assert template_csv().splitlines()[0] == ",".join(c.heading for c in COLUMNS)
        workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))
        assert workbook.sheetnames == ["Routes", "Notes", "Lists"]
        lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
        assert lists[0] == ["Area code", "Area name", "Route type code", "Route type"]
        assert lists[1] == ["SOUTH", "South", "URBAN", "Urban beat"]
        assert [row[:2] for row in lists[2:]] == [["CITY", "City"], ["RURAL", "Rural"]]


def test_days_read_as_names_or_numbers() -> None:
    """Mon, monday and 1 are the same day; an unknown token fails the cell."""
    assert parse_days("Mon, thursday") == [1, 4]
    assert parse_days("1;4 7") == [1, 4, 7]
    assert parse_days("Mon, Funday") is None
    assert parse_days("8") is None


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        report = _run(
            session,
            firm.id,
            _csv(
                "Beat Code,Beat Name,Area,Type,Frequency,Days,Status",
                "B01,Market Road,CITY,URBAN,Weekly,Mon,",
                "B02,,CITY,,,,",
                "B03,Temple St,NOWHERE,,,,",
                "B04,Bus Stand,CITY,NOPE,Hourly,Funday,Gone",
                "b01,Again,CITY,,,,",
                ",Blank,CITY,,,,",
                "CITY,Area clash,RURAL,,,,",
                "B05,Village,Rural,urban beat,on demand,1;4,inactive",
            ),
            apply=False,
        )
    found = {(issue.row, issue.column) for issue in report.issues}
    assert (3, "Name") in found
    assert (4, "Parent") in found
    assert {(5, "Type"), (5, "Frequency"), (5, "Days"), (5, "Status")} <= found
    assert (6, "Code") in found  # appears again
    assert (7, "Code") in found  # is required
    assert (8, "Code") in found  # already a territory's code
    assert len(report.issues) == 9
    assert report.to_create == 2  # B01 and B05 staged, then rolled back
    assert report.imported is False
    assert _routes(factory) == []


def test_an_apply_is_all_or_nothing() -> None:
    """One bad row keeps the good ones out; a clean file goes in whole."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        bad = _run(
            session,
            firm.id,
            _csv("Code,Name,Parent", "B01,Market Road,CITY", "B02,Temple,NOWHERE"),
            apply=True,
        )
        assert bad.imported is False
        assert _routes(factory) == []

        good = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Parent,Type,Frequency,Days,Status",
                'B01,Market Road,CITY,URBAN,Weekly,"Mon, Thu",',
                "B02,Village,Rural,,,,Inactive",
            ),
            apply=True,
        )
    assert good.imported is True
    assert good.to_create == 2
    assert _routes(factory) == ["B01", "B02"]
    with factory() as fresh:
        b01 = fresh.scalars(
            select(SalesTerritoryNode).where(SalesTerritoryNode.code == "B01")
        ).one()
        city = fresh.scalars(
            select(SalesTerritoryNode).where(SalesTerritoryNode.code == "CITY")
        ).one()
        assert b01.parent_id == city.id
        assert b01.path == "SOUTH/CITY/B01"
        assert b01.status == "ACTIVE"
        profile = fresh.scalars(
            select(TerritoryRouteProfile).where(
                TerritoryRouteProfile.territory_id == b01.id
            )
        ).one()
        urban = fresh.scalars(select(RouteTypeMaster)).one()
        assert profile.route_type_id == urban.id
        assert profile.visit_frequency == "WEEKLY"
        days = fresh.scalars(
            select(TerritoryWorkingDay.weekday).where(
                TerritoryWorkingDay.route_profile_id == profile.id
            )
        ).all()
        assert sorted(days) == [1, 4]
        b02 = fresh.scalars(
            select(SalesTerritoryNode).where(SalesTerritoryNode.code == "B02")
        ).one()
        assert b02.status == "INACTIVE"
        # Each route is audited as the form audits it.
        assert (
            len(
                fresh.scalars(
                    select(AuditLog).where(
                        AuditLog.action == "sales_territory.created",
                        AuditLog.entity_id.in_([b01.id, b02.id]),
                    )
                ).all()
            )
            == 2
        )


def test_a_second_apply_refuses_the_same_codes_unless_asked_to_update() -> None:
    """A code already a route is refused; update changes it, keeping blanks."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        first = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Parent,Type,Frequency,Days",
                "B01,Market Road,CITY,URBAN,Weekly,Mon",
            ),
            apply=True,
        )
        assert first.imported is True

        again = _run(session, firm.id, _csv("Code,Name", "b01,Other"), apply=True)
        assert again.imported is False
        assert [(i.row, i.column) for i in again.issues] == [(2, "Code")]
        assert "is already a route" in again.issues[0].message

        updated = _run(
            session,
            firm.id,
            _csv("Code,Name,Parent,Days,Status", "B01,,RURAL,Tue,Inactive"),
            apply=True,
            existing="update",
        )
        assert updated.imported is True, updated.issues
        assert updated.to_update == 1

        area = _run(
            session,
            firm.id,
            _csv("Code,Name", "CITY,City Beat"),
            apply=True,
            existing="update",
        )
        assert area.imported is False
        assert "is an area, not a route" in area.issues[0].message
    with factory() as fresh:
        b01 = fresh.scalars(
            select(SalesTerritoryNode).where(SalesTerritoryNode.code == "B01")
        ).one()
        # A blank Name keeps the one it had; the parent moved, the path with it.
        assert b01.name == "Market Road"
        assert b01.path == "SOUTH/RURAL/B01"
        assert b01.status == "INACTIVE"
        profile = fresh.scalars(
            select(TerritoryRouteProfile).where(
                TerritoryRouteProfile.territory_id == b01.id
            )
        ).one()
        # Type and frequency were left blank, so they stand.
        assert profile.visit_frequency == "WEEKLY"
        assert profile.route_type_id is not None
        days = fresh.scalars(
            select(TerritoryWorkingDay.weekday).where(
                TerritoryWorkingDay.route_profile_id == profile.id,
                TerritoryWorkingDay.is_deleted.is_(False),
            )
        ).all()
        assert list(days) == [2]


def test_routes_can_be_mapped_from_another_softwares_file() -> None:
    """The kind is registered under the route form's code, with these columns."""
    factory = _factory()
    assert IMPORT_KINDS["routes"].permission == "TERRITORY_CREATE"
    with factory() as session:
        assert columns_for_kind(session, "routes") == COLUMNS
