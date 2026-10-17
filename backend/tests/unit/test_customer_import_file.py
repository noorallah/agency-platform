"""Customers come in from a file, with a template to fill in (backlog 46).

The file rules are the product import's, from ``app.common.file_import``: every
row checked before anything is written, all or nothing, update by code. What
these tests pin is what is particular to customers -- one address and one
contact per row, updated in place; an opening balance booked to the ledger and
read the way Tally writes it; a standing discount held to the settings code;
and a 10-digit phone number taken as Indian.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport, indian_phone
from app.core.database.base import Base
from app.customers.models import (
    Customer,
    CustomerAddress,
    CustomerContact,
    CustomerGroup,
    CustomerReceivableTransaction,
)
from app.customers.schemas import CustomerCreate
from app.customers.services import CustomerService
from app.customers.services.customer_import import (
    COLUMNS,
    CustomerFileImporter,
    template_csv,
    template_workbook,
)
from app.finance.models import JournalEntry, JournalLine, LedgerAccount
from app.finance.services.opening_setup import seed_finance_setup
from app.firms.models import Firm
from app.identity.models import User, UserFirm
from app.sales.models import (
    SalesTerritoryNode,
    TerritoryCustomerAssignment,
    TerritoryRouteProfile,
)
from app.salesmen.models import Salesman
from app.trade_licences.models import TradeLicence, TradeLicenceType


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session, *, books: bool = True) -> Firm:
    """Add a firm with one segment and, unless told not to, its books."""
    firm = Firm(
        name="Import Firm",
        code="IMPORT",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    session.add(CustomerGroup(firm_id=firm.id, code="RETAIL", name="Retail shops"))
    session.commit()
    if books:
        seed_finance_setup(
            session,
            firm_id=firm.id,
            year_starts_on=date(2026, 4, 1),
            actor_id=uuid4(),
        )
        session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
    file_format: str = "csv",
    may_manage_settings: bool = True,
    may_assign_routes: bool = True,
    may_record_licences: bool = True,
) -> ImportReport[Customer]:
    """Check or apply one file as a single actor."""
    return CustomerFileImporter(
        session,
        CustomerService(session),
        may_manage_settings=may_manage_settings,
        may_assign_routes=may_assign_routes,
        may_record_licences=may_record_licences,
    ).run(
        content,
        file_format=file_format,  # type: ignore[arg-type]
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(fresh.scalars(select(Customer.code)).all())


def _receivable_net(session: Session, firm_id: UUID) -> Decimal:
    """Return what the ledger says customers owe: the receivable's net debit."""
    rows = session.execute(
        select(JournalLine.debit_amount, JournalLine.credit_amount)
        .join(LedgerAccount, LedgerAccount.id == JournalLine.ledger_account_id)
        .join(JournalEntry, JournalEntry.id == JournalLine.journal_entry_id)
        .where(LedgerAccount.firm_id == firm_id, LedgerAccount.code == "1100")
    ).all()
    return sum((Decimal(debit) - Decimal(credit) for debit, credit in rows), Decimal(0))


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _run(session, firm.id, _csv("Code,Name", "TAKEN,Already here"), apply=True)
    content = _csv(
        "Code,Name,Segment,Phone,Email,CreditLimit,Type,City,OpeningBalance",
        "GOOD-1,Good,Retail shops,9876543210,a@b.example,1000,,,",
        "BAD-SEG,Bad segment,WHOLESALE,,,,,,",
        "BAD-PHONE,Bad phone,,12345,,,,,",
        "BAD-MAIL,Bad email,,,not-an-email,,,,",
        "BAD-NUM,Bad number,,,,lots,,,",
        "BAD-TYPE,Bad type,,,,,COMPANY,,",
        "BAD-ADDR,Half an address,,,,,,Chennai,",
        "BAD-OPEN,Bad balance,,,,,,,twelve Dr",
        "TAKEN,Clash,,,,,,,",
        "GOOD-1,Twice,,,,,,,",
        "NO-NAME,,,,,,,,",
    )

    report = _run(session, firm.id, content, apply=True)

    found = {(issue.row, issue.column) for issue in report.issues}
    assert found == {
        (3, "Segment"),
        (4, "Phone"),
        (5, "Email"),
        (6, "CreditLimit"),
        (7, "Type"),
        (8, "Address1"),
        (8, "State"),
        (8, "PIN"),
        (9, "OpeningBalance"),
        (10, "Code"),
        (11, "Code"),
        (12, "Name"),
    }, [issue.describe() for issue in report.issues]
    assert not report.imported
    assert _codes(factory) == ["TAKEN"]


def test_a_clean_file_imports_whole_with_its_address_contact_and_balance() -> None:
    """Segment, address, contact and a Tally-style balance all land."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Party Code,Party Name,Customer Group,GSTIN/UIN,Mobile No,Credit Days,"
        "Opening Balance,Address,City,State,Pincode,Contact Person,Contact Mobile",
        'OWES,Owes Us,RETAIL,33AAAPL1234C1Z5,098765 43210,30,"1,200.50 Dr",'
        "12 Market Road,Chennai,Tamil Nadu,600001,R. Kumar,9123456789",
        "PAID,Paid Ahead,,,,,500 Cr,,,,,,",
        "PLAIN,Plain,,,,,,,,,,,",
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    assert (report.to_create, report.to_update) == (3, 0)
    with factory() as fresh:
        owes = fresh.scalars(select(Customer).where(Customer.code == "OWES")).one()
        assert owes.customer_group_id is not None
        assert owes.phone == "+919876543210"
        assert owes.currency_code == "INR"
        assert owes.customer_type == "BUSINESS"
        assert owes.payment_terms_days == 30
        assert owes.opening_balance == Decimal("1200.50")
        assert owes.current_outstanding == Decimal("1200.50")
        (address,) = owes.addresses
        assert (address.city, address.postal_code, address.country) == (
            "Chennai",
            "600001",
            "IN",
        )
        assert address.is_default_billing and address.is_default_shipping
        (contact,) = owes.contacts
        assert (contact.name, contact.mobile, contact.is_primary) == (
            "R. Kumar",
            "+919123456789",
            True,
        )
        paid = fresh.scalars(select(Customer).where(Customer.code == "PAID")).one()
        assert paid.opening_balance == Decimal("-500")
        assert paid.unapplied_advance_balance == Decimal("500")
        # Booked, not just recorded: the receivable carries the net of both.
        assert _receivable_net(fresh, firm.id) == Decimal("700.50")
        assert len(fresh.scalars(select(CustomerReceivableTransaction)).all()) == 2
        created = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "customer.created")
        ).all()
        assert len(created) == 3
    assert _codes(factory) == ["OWES", "PAID", "PLAIN"]


def test_a_firm_without_books_is_told_per_row_and_nothing_is_written() -> None:
    """An opening balance has to be booked, so no chart means no balance."""
    factory = _factory()
    session = factory()
    firm = _firm(session, books=False)
    content = _csv("Code,Name,OpeningBalance", "NOBAL,No balance,", "BAL,Owes,100")

    report = _run(session, firm.id, content, apply=True)

    assert [(issue.row, issue.code) for issue in report.issues] == [(3, "BAL")]
    assert "chart of accounts" in report.issues[0].message
    assert _codes(factory) == []


def test_updating_by_code_changes_only_what_the_file_fills() -> None:
    """A blank cell keeps its field; the address changes in place."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    CustomerService(session).create(
        CustomerCreate.model_validate(
            {
                "code": "KEEP",
                "customer_type": "BUSINESS",
                "name": "Keep Me",
                "email": "keep@shop.example",
                "currency_code": "INR",
                "addresses": [
                    {
                        "address_type": "BILLING",
                        "address_line1": "1 Old Street",
                        "city": "Madurai",
                        "state": "Tamil Nadu",
                        "country": "IN",
                        "postal_code": "625001",
                        "is_default_billing": True,
                    },
                    {
                        "address_type": "SHIPPING",
                        "address_line1": "Godown 4",
                        "city": "Madurai",
                        "state": "Tamil Nadu",
                        "country": "IN",
                        "postal_code": "625002",
                        "is_default_shipping": True,
                    },
                ],
            }
        ),
        firm_id=firm.id,
        actor_id=uuid4(),
    )
    content = _csv(
        "Code,Name,Email,Phone,City,PIN",
        "KEEP,,,9876543210,Chennai,600001",
        "NEW-1,Brand new,,,,",
    )

    refused = _run(session, firm.id, content, apply=True)
    assert [issue.row for issue in refused.issues] == [2]
    assert "update existing customers" in refused.issues[0].message

    report = _run(session, firm.id, content, apply=True, existing="update")

    assert report.imported, [issue.describe() for issue in report.issues]
    assert (report.to_create, report.to_update) == (1, 1)
    with factory() as fresh:
        kept = fresh.scalars(select(Customer).where(Customer.code == "KEEP")).one()
        assert (kept.name, kept.email, kept.phone) == (
            "Keep Me",
            "keep@shop.example",
            "+919876543210",
        )
        addresses = fresh.scalars(
            select(CustomerAddress)
            .where(CustomerAddress.customer_id == kept.id)
            .order_by(CustomerAddress.postal_code)
        ).all()
        assert [
            (a.address_line1, a.city, a.postal_code, a.is_deleted) for a in addresses
        ] == [
            ("1 Old Street", "Chennai", "600001", False),
            ("Godown 4", "Madurai", "625002", False),
        ]
        assert fresh.scalars(select(CustomerContact)).all() == []
    assert _codes(factory) == ["KEEP", "NEW-1"]


def test_a_standing_discount_needs_the_settings_permission() -> None:
    """A file is a second way to set a price, and answers to the same code."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv("Code,Name,DiscountPercent", "DISC,Discounted,5", "NONE,Plain,")

    refused = _run(session, firm.id, content, apply=True, may_manage_settings=False)

    assert [(issue.row, issue.code) for issue in refused.issues] == [(2, "DISC")]
    assert "CUSTOMER_MANAGE_SETTINGS" in refused.issues[0].message
    assert _codes(factory) == []

    allowed = _run(session, firm.id, content, apply=True)
    assert allowed.imported
    assert _codes(factory) == ["DISC", "NONE"]


def test_one_gst_number_on_two_rows_is_refused_and_names_the_first() -> None:
    """A GST number is held once (backlog 125); the second row is a shop."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Code,Name,GSTIN",
        "FIRST,First,33AAAPL1234C1Z5",
        "SECOND,Second,33aaapl1234c1z5",
    )

    report = _run(session, firm.id, content, apply=True)

    assert [issue.code for issue in report.issues] == ["SECOND"]
    assert "Make this customer a shop of FIRST" in report.issues[0].message
    assert _codes(factory) == []


def test_a_bare_indian_number_is_written_with_its_country_code() -> None:
    """Ten digits, a trunk 0 or a bare 91 all mean India; + is left alone."""
    assert indian_phone("98765 43210") == "+919876543210"
    assert indian_phone("09876543210") == "+919876543210"
    assert indian_phone("919876543210") == "+919876543210"
    assert indian_phone("+44 20 7946 0000") == "+442079460000"


def test_the_template_round_trips_and_lists_the_firms_segments() -> None:
    """The template's own example row imports, and its lists are this firm's."""
    factory = _factory()
    session = factory()
    firm = _firm(session)

    content = template_workbook(session, firm.id)

    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Customers", "Notes", "Lists"]
    headings = [cell.value for cell in workbook["Customers"][1]]
    assert headings == [column.heading for column in COLUMNS]
    listed = {
        value
        for row in workbook["Lists"].iter_rows(values_only=True)
        for value in row
        if value
    }
    assert {"RETAIL", "Retail shops", "BUSINESS", "ON_HOLD"} <= listed
    assert template_csv().splitlines()[0].split(",")[0] == "Code"

    report = _run(session, firm.id, content, apply=True, file_format="xlsx")
    assert report.issues == [], [issue.describe() for issue in report.issues]
    assert _codes(factory) == ["SHOP-001"]


def test_an_account_manager_is_named_from_the_salesman_list_not_the_members() -> None:
    """A member's email names nobody; the list's code or name names the entry.

    The member fallback went with the bridge: the firm's list of salesmen is
    the one authority on who may look after a customer.
    """
    factory = _factory()
    session = factory()
    firm = _firm(session)
    user = User(email="ravi@import.example", full_name="Ravi Kumar", password_hash="h")
    session.add(user)
    session.flush()
    session.add(UserFirm(user_id=user.id, firm_id=firm.id, is_active=True))
    entry = Salesman(firm_id=firm.id, code="RK", name="Ravi K", user_id=user.id)
    session.add(entry)
    session.commit()

    for given in ("ravi@import.example", "Ravi Kumar"):
        refused = _run(
            session,
            firm.id,
            _csv("Code,Name,AccountManager", f"SHOP-1,Shop,{given}"),
            apply=True,
        )
        assert [(issue.column, issue.message) for issue in refused.issues] == [
            ("AccountManager", f"'{given}' is not on the firm's list of salesmen.")
        ]
    assert _codes(factory) == []

    taken = _run(
        session,
        firm.id,
        _csv("Code,Name,AccountManager", "SHOP-1,Shop,rk"),
        apply=True,
    )
    assert taken.issues == [], [issue.describe() for issue in taken.issues]
    with factory() as fresh:
        stored = fresh.scalar(select(Customer).where(Customer.code == "SHOP-1"))
        assert stored is not None
        assert stored.salesman_id == entry.id


def _route(
    session: Session, firm_id: UUID, code: str, name: str, *, area: bool = False
) -> SalesTerritoryNode:
    """Add a territory node; a route unless ``area``, which has no profile."""
    node = SalesTerritoryNode(
        firm_id=firm_id, hierarchy_level_id=uuid4(), code=code, name=name, path=code
    )
    session.add(node)
    session.flush()
    if not area:
        session.add(TerritoryRouteProfile(territory_id=node.id))
    session.commit()
    return node


def _on_route(session: Session, route_id: UUID) -> dict[str, int | None]:
    """Return the live customers on a route by code, with their stop."""
    rows = session.execute(
        select(Customer.code, TerritoryCustomerAssignment.visit_sequence)
        .join(Customer, Customer.id == TerritoryCustomerAssignment.customer_id)
        .where(
            TerritoryCustomerAssignment.territory_id == route_id,
            TerritoryCustomerAssignment.is_deleted.is_(False),
        )
    ).all()
    return {code: stop for code, stop in rows}


def test_a_beat_puts_the_customer_on_an_existing_route_and_keeps_its_round() -> None:
    """By code or name; the route's other shops and their stops stay put (G5)."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    beat = _route(session, firm.id, "BEAT-01", "Market Road")
    _run(session, firm.id, _csv("Code,Name", "FIRST,First shop"), apply=True)
    first = session.scalars(select(Customer).where(Customer.code == "FIRST")).one()
    session.add(
        TerritoryCustomerAssignment(
            territory_id=beat.id, customer_id=first.id, visit_sequence=1
        )
    )
    session.commit()

    report = _run(
        session,
        firm.id,
        _csv("Code,Name,Route", "BYCODE,By code,beat-01", "BYNAME,By name,market road"),
        apply=True,
    )

    assert report.imported, [issue.describe() for issue in report.issues]
    with factory() as fresh:
        assert _on_route(fresh, beat.id) == {"FIRST": 1, "BYCODE": None, "BYNAME": None}
        audits = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "sales_territory.customers_set")
        ).all()
        assert len(audits) == 2

    again = _run(
        session,
        firm.id,
        _csv("Code,Name,Beat", "BYCODE,,BEAT-01"),
        apply=True,
        existing="update",
    )
    assert again.imported, [issue.describe() for issue in again.issues]
    with factory() as fresh:
        assert len(_on_route(fresh, beat.id)) == 3


def test_an_unknown_beat_or_an_area_is_refused_and_nothing_is_written() -> None:
    """The file never makes a route, nor hangs a shop off an area."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _route(session, firm.id, "SOUTH", "South Zone", area=True)
    content = _csv(
        "Code,Name,Beat", "NOWHERE,Nowhere,BEAT-99", "AREA,Area,SOUTH", "NONE,None,"
    )

    report = _run(session, firm.id, content, apply=True)

    assert [(issue.row, issue.column) for issue in report.issues] == [
        (2, "Beat"),
        (3, "Beat"),
    ]
    assert "'BEAT-99' is not one of the firm's routes" in report.issues[0].message
    assert _codes(factory) == []


def test_a_beat_needs_the_permission_the_route_screen_needs() -> None:
    """A file is a second way to fill a round, and answers to the same code."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _route(session, firm.id, "BEAT-01", "Market Road")

    refused = _run(
        session,
        firm.id,
        _csv("Code,Name,Beat", "SHOP,Shop,BEAT-01"),
        apply=True,
        may_assign_routes=False,
    )

    assert [issue.column for issue in refused.issues] == ["Beat"]
    assert "TERRITORY_ASSIGN_CUSTOMERS" in refused.issues[0].message
    assert _codes(factory) == []


def _licences(session: Session) -> list[tuple[str, str, date | None]]:
    """Return every live licence as (customer code, number, valid to)."""
    rows = session.execute(
        select(Customer.code, TradeLicence.licence_number, TradeLicence.valid_to)
        .join(Customer, Customer.id == TradeLicence.customer_id)
        .where(TradeLicence.is_deleted.is_(False))
        .order_by(Customer.code, TradeLicence.licence_number)
    ).all()
    return [(code, number, valid_to) for code, number, valid_to in rows]


def test_a_licence_goes_on_the_register_and_is_renewed_by_its_number() -> None:
    """Type, number and expiry land on trade_licences; a re-import renews (G5)."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Code,Name,Licence Type,DL No,Licence Expiry",
        "PHARMA,Pharma,DRUG_WHOLESALE,tn-01-20b-123,31-03-2027",
        "OTHER,Other,Other licence,LIC-9,",
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    with factory() as fresh:
        assert _licences(fresh) == [
            ("OTHER", "LIC-9", None),
            ("PHARMA", "TN-01-20B-123", date(2027, 3, 31)),
        ]
        # The usual types were seeded, as the licence screen's first read does.
        assert len(fresh.scalars(select(TradeLicenceType)).all()) >= 8
        created = fresh.scalars(
            select(AuditLog).where(AuditLog.action == "trade_licence.created")
        ).all()
        assert len(created) == 2

    renewed = _run(
        session,
        firm.id,
        _csv(
            "Code,Name,LicenceType,LicenceNumber,LicenceExpiry",
            "PHARMA,,DRUG_WHOLESALE,TN-01-20B-123,2028-03-31",
            "OTHER,,OTHER,LIC-10,",
        ),
        apply=True,
        existing="update",
    )

    assert renewed.imported, [issue.describe() for issue in renewed.issues]
    with factory() as fresh:
        # Renewed in place; a new number is a new licence, the old one kept.
        assert _licences(fresh) == [
            ("OTHER", "LIC-10", None),
            ("OTHER", "LIC-9", None),
            ("PHARMA", "TN-01-20B-123", date(2028, 3, 31)),
        ]


def test_a_drug_licence_with_no_expiry_loads_blank() -> None:
    """A number with no expiry passes the check and is loaded with none.

    Owner, 2026-10-17: an old file often has the licence number and not the
    date; the row is not an error, and no date is invented for it. Every
    other optional cell left blank is no error either.
    """
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Code,Name,Address1,City,PIN,Phone,Email,PAN,GSTIN,"
        "LicenceType,LicenceNumber,LicenceExpiry",
        "NOEXP,No expiry,,,,,,,,DRUG_WHOLESALE,dl-1,",
    )

    checked = _run(session, firm.id, content, apply=False)
    assert not checked.issues, [issue.describe() for issue in checked.issues]

    report = _run(session, firm.id, content, apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    with factory() as fresh:
        assert _licences(fresh) == [("NOEXP", "DL-1", None)]


def test_a_licence_row_is_refused_by_the_cell_that_is_wrong() -> None:
    """Each refusal names its column, and the file writes nothing."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "Code,Name,LicenceType,LicenceNumber,LicenceExpiry",
        "BADTYPE,Bad type,LIQUOR,DL-2,31-03-2027",
        "NOTYPE,No type,,DL-3,31-03-2027",
        "BADDATE,Bad date,FSSAI,DL-4,someday",
        "NONUM,No number,FSSAI,,31-03-2027",
    )

    report = _run(session, firm.id, content, apply=True)

    assert [(issue.row, issue.column) for issue in report.issues] == [
        (2, "LicenceType"),
        (3, "LicenceType"),
        (4, "LicenceExpiry"),
        (5, "LicenceNumber"),
    ], [issue.describe() for issue in report.issues]
    assert _codes(factory) == []
    with factory() as fresh:
        assert _licences(fresh) == []

    refused = _run(
        session,
        firm.id,
        _csv("Code,Name,LicenceType,LicenceNumber", "SHOP,Shop,OTHER,DL-5"),
        apply=True,
        may_record_licences=False,
    )
    assert [issue.column for issue in refused.issues] == ["LicenceNumber"]
    assert "TRADE_LICENCE_MANAGE" in refused.issues[0].message


def test_the_template_lists_the_firms_routes_and_licence_types() -> None:
    """A Beat and a LicenceType are picked from the Lists sheet."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    _route(session, firm.id, "BEAT-01", "Market Road")

    workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))

    listed = {
        value
        for row in workbook["Lists"].iter_rows(values_only=True)
        for value in row
        if value
    }
    assert {"BEAT-01", "Market Road", "DRUG_WHOLESALE", "FSSAI"} <= listed
