"""Customer segments (outlet types) come in from a file (Swathi G3).

The file rules are the other masters', from ``app.common.file_import``. What
these tests pin is what is particular to segments -- a row is keyed on its
code, a price level must already exist, and every row goes through the
segment service's own guards before the file commits once.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.customers.models import CustomerGroup
from app.customers.schemas import CustomerGroupWrite
from app.customers.services import CustomerGroupService
from app.customers.services.segment_import import (
    COLUMNS,
    SegmentFileImporter,
    template_csv,
    template_workbook,
)
from app.firms.models import Firm
from app.imports.services import IMPORT_KINDS, columns_for_kind
from app.pricing.models import PriceLevel

_ACTOR = uuid4()


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with one price level and one segment, Retailer."""
    firm = Firm(
        name="Segment Firm",
        code="SEGS",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.commit()
    session.add(PriceLevel(firm_id=firm.id, code="STD", name="Standard"))
    session.commit()
    CustomerGroupService(session).create_group(
        CustomerGroupWrite(code="RETAIL", name="Retailer"),
        firm_id=firm.id,
        actor_id=_ACTOR,
    )
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
) -> ImportReport[CustomerGroup]:
    """Check or apply one file as a single actor."""
    return SegmentFileImporter(session, CustomerGroupService(session)).run(
        content,
        file_format="csv",
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str]:
    """Read the codes of the live segments, from a session of its own."""
    with factory() as fresh:
        return sorted(
            fresh.scalars(
                select(CustomerGroup.code).where(CustomerGroup.is_deleted.is_(False))
            ).all()
        )


def _segment(factory: sessionmaker[Session], code: str) -> CustomerGroup:
    """Read one segment by code, from a session of its own."""
    with factory() as fresh:
        return fresh.scalars(
            select(CustomerGroup).where(CustomerGroup.code == code)
        ).one()


def test_the_template_names_the_columns_and_lists_segments_and_levels() -> None:
    """The CSV carries the headings; the workbook lists segments and levels."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        assert template_csv().splitlines()[0] == ",".join(c.heading for c in COLUMNS)
        workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))
        assert workbook.sheetnames == ["Segments", "Notes", "Lists"]
        lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
        assert lists[0] == [
            "Segment code",
            "Segment name",
            "Price level code",
            "Price level",
        ]
        assert lists[1] == ["RETAIL", "Retailer", "STD", "Standard"]


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        report = _run(
            session,
            firm.id,
            _csv(
                "Segment Code,Segment Name,Discount,Price Level,Active",
                "W01,Wholesaler,5,STD,Yes",
                "W02,,,,",
                "W03,Institution,150,,",
                "W04,Chemist,,NOPE,",
                "W05,Bakery,,,Maybe",
                "w01,Again,,,",
                ",Blank,,,",
                "RETAIL,Other,,,",
                "W06,Wholesaler,,,",
            ),
            apply=False,
        )
    found = {(issue.row, issue.column) for issue in report.issues}
    assert (3, "Name") in found
    assert (4, "Discount") in found
    assert (5, "PriceLevel") in found
    assert (6, "Active") in found
    assert (7, "Code") in found  # appears again
    assert (8, "Code") in found  # is required
    assert (9, "Code") in found  # already a segment
    # The name clash with W01, staged above it, is the service's own refusal.
    assert any(issue.row == 10 for issue in report.issues)
    assert len(report.issues) == 8
    assert report.to_create == 1  # W01 staged, then rolled back
    assert report.imported is False
    assert _codes(factory) == ["RETAIL"]


def test_an_apply_is_all_or_nothing() -> None:
    """One bad row keeps the good ones out; a clean file goes in whole."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        bad = _run(
            session,
            firm.id,
            _csv("Code,Name,PriceLevel", "W01,Wholesaler,STD", "W02,Chemist,NOPE"),
            apply=True,
        )
        assert bad.imported is False
        assert _codes(factory) == ["RETAIL"]

        good = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Description,Discount,PriceLevel,Active",
                "w01,Wholesaler,Bulk buyers,2.5,Standard,",
                "W02,Chemist,,,,No",
            ),
            apply=True,
        )
    assert good.imported is True, good.issues
    assert good.to_create == 2
    assert _codes(factory) == ["RETAIL", "W01", "W02"]
    w01 = _segment(factory, "W01")
    assert w01.name == "Wholesaler"
    assert w01.description == "Bulk buyers"
    assert w01.default_discount_percent == Decimal("2.5")
    assert w01.price_level_id is not None
    assert w01.is_active is True
    w02 = _segment(factory, "W02")
    assert w02.default_discount_percent == Decimal("0")
    assert w02.price_level_id is None
    assert w02.is_active is False
    with factory() as fresh:
        # Each segment is audited as the form audits it.
        audited = fresh.scalars(
            select(AuditLog).where(
                AuditLog.action == "customer_group.created",
                AuditLog.entity_id.in_([w01.id, w02.id]),
            )
        ).all()
        assert len(audited) == 2


def test_a_second_apply_refuses_the_same_codes_unless_asked_to_update() -> None:
    """A code already a segment is refused; update changes it, keeping blanks."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        first = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Description,Discount,PriceLevel",
                "W01,Wholesaler,Bulk buyers,5,STD",
            ),
            apply=True,
        )
        assert first.imported is True

        again = _run(session, firm.id, _csv("Code,Name", "w01,Other"), apply=True)
        assert again.imported is False
        assert [(i.row, i.column) for i in again.issues] == [(2, "Code")]
        assert "is already a segment" in again.issues[0].message

        updated = _run(
            session,
            firm.id,
            _csv("Code,Name,Discount,Active", "W01,,7.5,No"),
            apply=True,
            existing="update",
        )
        assert updated.imported is True, updated.issues
        assert updated.to_update == 1
    w01 = _segment(factory, "W01")
    # Blank or absent cells keep what the segment held.
    assert w01.name == "Wholesaler"
    assert w01.description == "Bulk buyers"
    assert w01.price_level_id is not None
    assert w01.default_discount_percent == Decimal("7.5")
    assert w01.is_active is False


def test_segments_can_be_mapped_from_another_softwares_file() -> None:
    """The kind is registered under the segment form's code, with these columns."""
    factory = _factory()
    assert IMPORT_KINDS["segments"].permission == "CUSTOMER_MANAGE_SETTINGS"
    with factory() as session:
        assert columns_for_kind(session, "segments") == COLUMNS
