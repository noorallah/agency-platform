"""Principals come in from a file, with a template to fill in (Swathi G3).

The file rules are the other masters', from ``app.common.file_import``. What
these tests pin is what is particular to principals -- the code is the key, a
supplier must already exist, and every row goes through the service's own
guards before the file commits once.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.common.audit.models import AuditLog
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.firms.models import Firm
from app.imports.services import IMPORT_KINDS, columns_for_kind
from app.products.models.brand import Principal
from app.products.services.brands import BrandService, PrincipalWrite
from app.products.services.principal_import import (
    COLUMNS,
    PrincipalFileImporter,
    template_csv,
    template_workbook,
)
from app.vendors.models import Vendor


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session) -> Firm:
    """Add a firm with one principal and one supplier."""
    firm = Firm(
        name="Principal Firm",
        code="PRINCS",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    session.add_all(
        [
            Principal(firm_id=firm.id, code="GENERAL", name="General"),
            Vendor(
                firm_id=firm.id,
                code="SUP01",
                name="Sri Traders",
                display_name="Sri Traders",
            ),
        ]
    )
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    existing: str = "refuse",
) -> ImportReport[Principal]:
    """Check or apply one file as a single actor."""
    return PrincipalFileImporter(session, BrandService(session)).run(
        content,
        file_format="csv",
        firm_id=firm_id,
        actor_id=uuid4(),
        existing=existing,  # type: ignore[arg-type]
        apply=apply,
    )


def _codes(factory: sessionmaker[Session]) -> list[str]:
    """Read what is durably stored, from a session of its own."""
    with factory() as fresh:
        return sorted(
            fresh.scalars(
                select(Principal.code).where(Principal.is_deleted.is_(False))
            ).all()
        )


def test_the_template_names_the_columns_and_lists_companies_and_suppliers() -> None:
    """The CSV carries the headings; the workbook lists what may be named."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        assert template_csv().splitlines()[0] == ",".join(c.heading for c in COLUMNS)
        workbook = load_workbook(BytesIO(template_workbook(session, firm.id)))
        assert workbook.sheetnames == ["Principals", "Notes", "Lists"]
        lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
        assert lists[0] == [
            "Company code",
            "Company name",
            "Supplier code",
            "Supplier name",
        ]
        assert lists[1] == ["GENERAL", "General", "SUP01", "Sri Traders"]


def test_a_check_reports_every_problem_by_row_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        report = _run(
            session,
            firm.id,
            _csv(
                "Company Code,Company Name,Vendor,Active",
                "HUL,Hindustan Unilever,SUP01,Yes",
                "ITC,,,",
                "NES,Nestle,NOBODY,",
                "DAB,Dabur,,maybe",
                "hul,Again,,",
                ",Nameless,,",
                "MAR,Marico,Sri Traders,No",
                "GENERAL,General Again,,",
            ),
            apply=False,
        )
    found = {(issue.row, issue.column, issue.message) for issue in report.issues}
    assert (3, "Name", "is needed for a new principal.") in found
    assert any(row == 4 and column == "Supplier" for row, column, _ in found)
    assert (5, "Active", "'maybe' should be Yes or No.") in found
    assert (6, "Code", "appears again; row 2 already has it.") in found
    assert (7, "Code", "is required.") in found
    assert any(row == 9 and column == "Code" for row, column, _ in found)
    assert len(report.issues) == 6
    assert report.to_create == 2  # HUL and MAR staged, then rolled back
    assert report.imported is False
    assert _codes(factory) == ["GENERAL"]


def test_an_apply_is_all_or_nothing() -> None:
    """One bad row keeps the good ones out; a clean file goes in whole."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        bad = _run(
            session,
            firm.id,
            _csv("Code,Name,Supplier", "HUL,Hindustan Unilever,", "ITC,ITC,NOBODY"),
            apply=True,
        )
        assert bad.imported is False
        assert _codes(factory) == ["GENERAL"]

        good = _run(
            session,
            firm.id,
            _csv(
                "Code,Name,Supplier,Active",
                "hul,Hindustan Unilever,sup01,",
                "MAR,Marico,,No",
            ),
            apply=True,
        )
    assert good.imported is True
    assert good.to_create == 2
    assert _codes(factory) == ["GENERAL", "HUL", "MAR"]
    with factory() as fresh:
        hul = fresh.scalars(select(Principal).where(Principal.code == "HUL")).one()
        marico = fresh.scalars(select(Principal).where(Principal.code == "MAR")).one()
        supplier = fresh.scalars(select(Vendor).where(Vendor.code == "SUP01")).one()
        assert hul.vendor_id == supplier.id
        assert hul.is_active is True
        assert marico.vendor_id is None
        assert marico.is_active is False
        # Each principal is audited as the form audits it.
        assert (
            len(
                fresh.scalars(
                    select(AuditLog).where(AuditLog.action == "principal.saved")
                ).all()
            )
            == 2
        )


def test_a_second_apply_refuses_the_same_codes_unless_asked_to_update() -> None:
    """A code already a principal is refused; update changes it."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        first = _run(
            session,
            firm.id,
            _csv("Code,Name,Supplier", "HUL,Hindustan Unilever,SUP01"),
            apply=True,
        )
        assert first.imported is True

        again = _run(session, firm.id, _csv("Code,Name", "hul,HUL Ltd"), apply=True)
        assert again.imported is False
        assert [(i.row, i.column) for i in again.issues] == [(2, "Code")]
        assert "is already a principal" in again.issues[0].message

        updated = _run(
            session,
            firm.id,
            _csv("Code,Name,Active", "HUL,,No"),
            apply=True,
            existing="update",
        )
    assert updated.imported is True
    assert updated.to_update == 1
    with factory() as fresh:
        hul = fresh.scalars(select(Principal).where(Principal.code == "HUL")).one()
        supplier = fresh.scalars(select(Vendor).where(Vendor.code == "SUP01")).one()
        # A blank cell or a missing column keeps what the principal held.
        assert hul.name == "Hindustan Unilever"
        assert hul.vendor_id == supplier.id
        assert hul.is_active is False


def test_save_principal_still_commits_on_its_own() -> None:
    """The form's path is the stage method and a commit, as before."""
    factory = _factory()
    with factory() as session:
        firm = _firm(session)
        saved = BrandService(session).save_principal(
            PrincipalWrite(code="itc", name="ITC Limited"),
            firm_id=firm.id,
            actor_id=uuid4(),
        )
    assert saved.code == "ITC"
    assert _codes(factory) == ["GENERAL", "ITC"]


def test_principals_can_be_mapped_from_another_softwares_file() -> None:
    """The kind is registered under the principal code, with its columns."""
    factory = _factory()
    assert IMPORT_KINDS["principals"].permission == "PRODUCT_UPDATE"
    with factory() as session:
        assert columns_for_kind(session, "principals") == COLUMNS
