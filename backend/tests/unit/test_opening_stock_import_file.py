"""Opening stock comes in from one file, with a template (backlog 36, 46).

A firm moving from Tally or Excel brings its stock on hand at cutover as one
spreadsheet: every problem is reported by row and column before anything is
written, rows are grouped into one document per warehouse, and every document
is created and posted in one transaction -- or none is.

The sessions do not autoflush, because a request's do not.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.batch_serial.models import BatchRecord
from app.branches.models import Branch, Warehouse
from app.common.file_import import ImportReport
from app.core.database.base import Base
from app.finance.models import GLPosting, JournalEntry, LedgerAccount
from app.finance.services.opening_setup import seed_finance_setup
from app.firms.models import Firm
from app.inventory.models import (
    InventoryRecord,
    OpeningStockBatch,
    StockLedgerEntry,
)
from app.inventory.schemas import OpeningStockBatchCreate
from app.inventory.services import InventoryService
from app.inventory.services.opening_stock_import import (
    OpeningStockFileImporter,
    parse_date,
    template_csv,
    template_workbook,
)
from app.products.models import Product

_ON = date(2026, 8, 1)


def _factory() -> sessionmaker[Session]:
    """Build one in-memory database whose sessions do not autoflush."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False, autoflush=False)


def _firm(session: Session, *, warehouses: tuple[str, ...] = ("MAIN", "EAST")) -> Firm:
    """Add a firm with its books, one branch, the warehouses and four products.

    RICE is plain; SYRUP is batch- and expiry-tracked; PHONE is serial-numbered;
    OIL is plain too, so a file can name two products.
    """
    actor = uuid4()
    firm = Firm(
        name="Stock Firm",
        code="STOCK",
        country="IN",
        currency_code="INR",
        financial_year_start=date(2026, 4, 1),
    )
    session.add(firm)
    session.flush()
    branch = Branch(
        firm_id=firm.id,
        code="HO",
        name="Head Office",
        display_name="Head Office",
        working_hours={},
        status="ACTIVE",
        created_by=actor,
        updated_by=actor,
    )
    session.add(branch)
    session.flush()
    for code in warehouses:
        session.add(
            Warehouse(
                firm_id=firm.id,
                branch_id=branch.id,
                code=code,
                name=f"{code.title()} Godown",
                display_name=f"{code.title()} Godown",
                status="ACTIVE",
                created_by=actor,
                updated_by=actor,
            )
        )
    for code, flags in (
        ("OIL", {}),
        ("PHONE", {"track_serial": True}),
        ("RICE", {}),
        ("SYRUP", {"track_batch": True, "track_expiry": True}),
    ):
        session.add(
            Product(
                firm_id=firm.id,
                code=code,
                name=code.title(),
                product_type="STOCK_ITEM",
                status="ACTIVE",
                created_by=actor,
                updated_by=actor,
                **flags,
            )
        )
    session.commit()
    seed_finance_setup(
        session, firm_id=firm.id, year_starts_on=date(2026, 4, 1), actor_id=actor
    )
    session.commit()
    return firm


def _csv(*lines: str) -> bytes:
    """Join lines into a CSV file's bytes."""
    return ("\n".join(lines) + "\n").encode("utf-8")


def _run(
    session: Session,
    firm_id: UUID,
    content: bytes,
    *,
    apply: bool,
    file_format: str = "csv",
) -> ImportReport[OpeningStockBatch]:
    """Check or apply one file as a single actor."""
    return OpeningStockFileImporter(session).run(
        content,
        file_format=file_format,  # type: ignore[arg-type]
        firm_id=firm_id,
        actor_id=uuid4(),
        posting_date=_ON,
        apply=apply,
    )


def _batches(factory: sessionmaker[Session]) -> list[str]:
    """Read the opening-stock documents durably stored, from a fresh session."""
    with factory() as fresh:
        return sorted(fresh.scalars(select(OpeningStockBatch.reference_number)).all())


def _open_warehouse(session: Session, firm: Firm, code: str) -> None:
    """Post opening stock for one warehouse the way the form does."""
    warehouse = session.scalars(select(Warehouse).where(Warehouse.code == code)).one()
    rice = session.scalars(select(Product).where(Product.code == "RICE")).one()
    service = InventoryService(session)
    actor = uuid4()
    batch = service.create_opening_stock_batch(
        OpeningStockBatchCreate(
            branch_id=warehouse.branch_id,
            warehouse_id=warehouse.id,
            reference_number="OS-BY-HAND",
            posting_date=_ON,
            lines=[{"product_id": rice.id, "quantity": "1"}],
        ),
        firm_id=firm.id,
        actor_id=actor,
    )
    service.post_opening_stock_batch(batch.id, firm_scope=firm.id, actor_id=actor)


def test_a_check_reports_every_problem_by_row_and_column_and_writes_nothing() -> None:
    """One pass names them all, with the column each belongs to."""
    factory = _factory()
    session = factory()
    firm = _firm(session, warehouses=("MAIN", "EAST", "OLD"))
    _open_warehouse(session, firm, "OLD")
    before = _batches(factory)
    content = _csv(
        "Item Code,Warehouse,Qty,Rate,Batch,Expiry",
        "RICE,MAIN,10,50,,",  # 2 fine
        "NOPE,MAIN,1,1,,",  # 3 unknown product
        "RICE,ATTIC,1,1,,",  # 4 unknown warehouse
        "OIL,MAIN,-2,1,,",  # 5 bad quantity
        "OIL,EAST,lots,1,,",  # 6 not a number
        "SYRUP,MAIN,5,10,,31-12-2027",  # 7 batch missing on a batch product
        "OIL,EAST,3,1,B1,",  # 8 batch on a plain product
        "SYRUP,EAST,5,10,B7,",  # 9 expiry missing
        "SYRUP,EAST,5,10,B8,31st Dec",  # 10 expiry unparseable
        "RICE,Main Godown,4,50,,",  # 11 duplicate of row 2, by name
        "RICE,OLD,4,50,,",  # 12 warehouse already opened
        "PHONE,MAIN,1,1000,,",  # 13 serial-numbered
        "OIL,,1,1,,",  # 14 warehouse blank with three of them
    )

    report = _run(session, firm.id, content, apply=False)

    found = {(issue.row, issue.column) for issue in report.issues}
    assert found == {
        (3, "ProductCode"),
        (4, "Warehouse"),
        (5, "Quantity"),
        (6, "Quantity"),
        (7, "Batch"),
        (8, "Batch"),
        (9, "Expiry"),
        (10, "Expiry"),
        (11, "ProductCode"),
        (12, "Warehouse"),
        (13, "ProductCode"),
        (14, "Warehouse"),
    }, [issue.describe() for issue in report.issues]
    messages = {issue.row: issue.describe() for issue in report.issues}
    assert "row 2 already has" in messages[11]
    assert "OS-BY-HAND" in messages[12] and "stock adjustment" in messages[12]
    assert "one serial at a time" in messages[13]
    assert report.rows == 13
    assert report.columns_used == [
        "ProductCode",
        "Warehouse",
        "Quantity",
        "UnitCost",
        "Batch",
        "Expiry",
    ]
    assert not report.imported
    assert _batches(factory) == before

    applied = _run(session, firm.id, content, apply=True)
    assert not applied.imported, "an apply with problems writes nothing either"
    assert _batches(factory) == before


def test_a_clean_file_creates_and_posts_one_document_per_warehouse() -> None:
    """Two warehouses, two posted documents; the stock is on the shelf and valued."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "ProductCode,Warehouse,Quantity,UnitCost,Batch,Expiry,Remarks",
        "RICE,MAIN,10,50,,,Counted on the day",
        "SYRUP,MAIN,4,25.50,B-MARCH,31/03/2027,",
        "SYRUP,MAIN,6,25.50,B-JUNE,2027-06-30,",
        "OIL,east godown,3,100,,,",
    )

    report = _run(session, firm.id, content, apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    assert report.to_create == 4
    assert _batches(factory) == [
        "OS-IMPORT-20260801-EAST",
        "OS-IMPORT-20260801-MAIN",
    ]
    with factory() as fresh:
        documents = fresh.scalars(select(OpeningStockBatch)).all()
        assert {row.status for row in documents} == {"POSTED"}
        assert {row.posting_date for row in documents} == {_ON}
        on_hand = {
            (product, warehouse, number): quantity
            for product, warehouse, number, quantity in fresh.execute(
                select(
                    Product.code,
                    Warehouse.code,
                    BatchRecord.batch_number,
                    InventoryRecord.current_quantity,
                )
                .join(Product, Product.id == InventoryRecord.product_id)
                .join(Warehouse, Warehouse.id == InventoryRecord.warehouse_id)
                .outerjoin(BatchRecord, BatchRecord.id == InventoryRecord.batch_id)
            ).all()
        }
        assert on_hand == {
            ("RICE", "MAIN", None): Decimal("10"),
            ("SYRUP", "MAIN", "B-MARCH"): Decimal("4"),
            ("SYRUP", "MAIN", "B-JUNE"): Decimal("6"),
            ("OIL", "EAST", None): Decimal("3"),
        }
        expiries = dict(
            fresh.execute(
                select(BatchRecord.batch_number, BatchRecord.expiry_date)
            ).all()
        )
        assert expiries == {
            "B-MARCH": date(2027, 3, 31),
            "B-JUNE": date(2027, 6, 30),
        }
        valued = fresh.scalar(select(func.sum(StockLedgerEntry.total_cost)))
        # 10 x 50 + 10 x 25.50 + 3 x 100
        assert Decimal(str(valued)) == Decimal("1055")
        postings = {
            reference: (debit, credit)
            for reference, debit, credit in fresh.execute(
                select(
                    JournalEntry.reference_number,
                    func.sum(GLPosting.debit_amount),
                    func.sum(GLPosting.credit_amount),
                )
                .join(JournalEntry, JournalEntry.id == GLPosting.journal_entry_id)
                .join(LedgerAccount, LedgerAccount.id == GLPosting.ledger_account_id)
                .where(LedgerAccount.code == "1200")
                .group_by(JournalEntry.reference_number)
            ).all()
        }
        assert postings == {
            "OS-IMPORT-20260801-MAIN": (Decimal("755.00"), Decimal("0.00")),
            "OS-IMPORT-20260801-EAST": (Decimal("300.00"), Decimal("0.00")),
        }


def test_one_bad_row_writes_nothing_at_all() -> None:
    """A clean warehouse is not written because another warehouse's row failed."""
    factory = _factory()
    session = factory()
    firm = _firm(session)
    content = _csv(
        "ProductCode,Warehouse,Quantity,UnitCost",
        "RICE,MAIN,10,50",
        "OIL,EAST,0,100",
    )

    report = _run(session, firm.id, content, apply=True)

    assert not report.imported
    assert [(issue.row, issue.column) for issue in report.issues] == [(3, "Quantity")]
    assert _batches(factory) == []
    with factory() as fresh:
        assert fresh.scalar(select(func.count()).select_from(InventoryRecord)) == 0


def test_the_only_warehouse_is_taken_when_the_column_is_blank() -> None:
    """A one-warehouse firm need not name it; a second import is refused by name."""
    factory = _factory()
    session = factory()
    firm = _firm(session, warehouses=("MAIN",))

    report = _run(session, firm.id, _csv("ProductCode,Quantity", "RICE,5"), apply=True)

    assert report.imported, [issue.describe() for issue in report.issues]
    assert _batches(factory) == ["OS-IMPORT-20260801-MAIN"]
    again = _run(session, firm.id, _csv("ProductCode,Quantity", "OIL,5"), apply=True)
    assert not again.imported
    assert "already has posted opening stock" in again.issues[0].message


def test_dates_are_read_the_ways_a_spreadsheet_writes_them() -> None:
    """Day first with - or /, ISO with or without a time, or an Excel serial."""
    assert parse_date("31-12-2027") == date(2027, 12, 31)
    assert parse_date("31/12/2027") == date(2027, 12, 31)
    assert parse_date("2027-12-31") == date(2027, 12, 31)
    assert parse_date("2027-12-31 00:00:00") == date(2027, 12, 31)
    assert parse_date("46752") == date(2027, 12, 31)
    assert parse_date("31-02-2027") is None
    assert parse_date("Dec 2027") is None


def test_the_template_round_trips_and_lists_the_firms_stock() -> None:
    """The example row names the firm's own product and warehouse, and imports."""
    factory = _factory()
    session = factory()
    firm = _firm(session)

    workbook_bytes = template_workbook(session, firm.id)
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)
    lists = [list(row) for row in workbook["Lists"].iter_rows(values_only=True)]
    assert lists[0][:5] == [
        "Warehouse code",
        "Warehouse name",
        None,
        "Product code",
        "Product name",
    ]
    listed = {row[3]: (row[5], row[6]) for row in lists[1:] if row[3]}
    assert "PHONE" not in listed, "serial-numbered stock is not imported"
    assert listed["SYRUP"] == ("Required", "Required")
    assert listed["RICE"] == (None, None)
    csv_text = template_csv(session, firm.id)
    assert csv_text.splitlines()[1].startswith("OIL,EAST,")

    report = _run(session, firm.id, workbook_bytes, apply=True, file_format="xlsx")

    assert report.imported, [issue.describe() for issue in report.issues]
    assert _batches(factory) == ["OS-IMPORT-20260801-EAST"]
