"""Build the customer-questions answer workbook from its markdown source.

``docs/CUSTOMER_REQUIREMENTS_QUESTIONS.md`` is the one list of questions. This
writes the same questions into a spreadsheet with cells for the customer's
answers and for the review afterwards, so the two cannot drift apart: change a
question in the markdown and run this again.

It writes a blank workbook and overwrites the output file, so a copy filled in
for a customer must be saved under another name.

    uv run python scripts/make_customer_questions_xlsx.py
    uv run python scripts/make_customer_questions_xlsx.py --output other.xlsx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPO_ROOT / "docs" / "CUSTOMER_REQUIREMENTS_QUESTIONS.md"
DEFAULT_OUTPUT = (
    REPO_ROOT / "dist" / "customer" / "CUSTOMER_REQUIREMENTS_QUESTIONS.xlsx"
)

FONT = "Arial"
OUTCOMES = '"Setting,Backlog item,Not done today,To check,No action"'
QUESTION_TITLES = [
    "#",
    "Section",
    "Ask",
    "What it decides",
    "Priority",
    "Check first?",
    "Customer's answer",
    "Outcome",
    "Review notes / what we do next",
    "Reviewed?",
]
QUESTION_WIDTHS = (7, 24, 46, 40, 10, 9, 50, 16, 40, 11)
PRIORITY_COLOURS = {"Critical": "C00000", "High": "C55A11", "Low": "595959"}
NOT_SHOWN_TITLES = [
    "What the customer does today",
    "Did we show it?",
    "Answer given",
    "Outcome",
    "Review notes / what we do next",
    "Reviewed?",
]
NOT_SHOWN_WIDTHS = (46, 14, 46, 16, 40, 11)
HEADER_ROW = 5
FIRST_ANSWER_COLUMN = 7
PRIORITY_COLUMN = 5
CHECK_COLUMN = 6
NOT_SHOWN_FIRST, NOT_SHOWN_LAST = 4, 23

_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
WRAP = Alignment(wrap_text=True, vertical="top")

Question = tuple[str, str, str, str, str, str]


def read_questions(source: Path) -> list[Question]:
    """Return (number, section, ask, decides, priority, check) per question."""
    questions: list[Question] = []
    section = ""
    for line in source.read_text(encoding="utf-8").splitlines():
        heading = re.match(r"## (\d+\. .+)", line)
        if heading:
            section = heading.group(1)
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if not re.fullmatch(r"\d+\.\d+", cells[0]):
            continue
        if len(cells) != 4 or cells[3] not in PRIORITY_COLOURS:
            raise SystemExit(
                f"Question {cells[0]} needs four cells ending in a priority "
                f"({', '.join(PRIORITY_COLOURS)})."
            )
        check = "Yes" if "**check**" in cells[2].lower() else ""
        number, ask, decides = (re.sub(r"\*+", "", cell) for cell in cells[:3])
        questions.append((number, section, ask, decides, cells[3], check))
    return questions


def write_header(sheet: Worksheet, row: int, titles: list[str]) -> None:
    """Write one styled header row."""
    for column, title in enumerate(titles, 1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def add_list(sheet: Worksheet, choices: str, cells: str) -> None:
    """Give a range of cells a drop-down of the named choices."""
    validation = DataValidation(type="list", formula1=choices, allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(cells)


def build_questions_sheet(sheet: Worksheet, questions: list[Question]) -> int:
    """Fill the Questions sheet and return its last row."""
    sheet.title = "Questions"
    for row, label in enumerate(("Customer", "Meeting date", "Met by"), 1):
        sheet.cell(row=row, column=1, value=label).font = Font(name=FONT, bold=True)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        for column in range(2, 5):
            sheet.cell(row=row, column=column).fill = INPUT_FILL
            sheet.cell(row=row, column=column).border = BORDER
        sheet.cell(row=row, column=2).font = Font(name=FONT)

    write_header(sheet, HEADER_ROW, QUESTION_TITLES)
    sheet.row_dimensions[HEADER_ROW].height = 30

    row = HEADER_ROW
    last_section = None
    for number, section, ask, decides, priority, check in questions:
        if section != last_section:
            row += 1
            title = sheet.cell(row=row, column=1, value=section)
            title.font = Font(name=FONT, bold=True)
            for column in range(1, len(QUESTION_TITLES) + 1):
                sheet.cell(row=row, column=column).fill = SECTION_FILL
            last_section = section
        row += 1
        values = [number, section, ask, decides, priority, check]
        values += [None] * (len(QUESTION_TITLES) - len(values))
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=column, value=value)
            is_check = column == CHECK_COLUMN
            colour = "C00000" if is_check else "000000"
            if column == PRIORITY_COLUMN:
                colour = PRIORITY_COLOURS[priority]
            cell.font = Font(
                name=FONT,
                bold=(is_check and bool(value)) or value == "Critical",
                color=colour,
            )
            cell.alignment = WRAP
            cell.border = BORDER
            if column >= FIRST_ANSWER_COLUMN:
                cell.fill = INPUT_FILL
        sheet.cell(row=row, column=1).number_format = "@"
        sheet.row_dimensions[row].height = 48
    last = row

    for letter, width in zip("ABCDEFGHIJ", QUESTION_WIDTHS, strict=True):
        sheet.column_dimensions[letter].width = width
    first = HEADER_ROW + 1
    sheet.freeze_panes = f"C{first}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:J{last}"
    add_list(sheet, OUTCOMES, f"H{first}:H{last}")
    add_list(sheet, '"Yes,No"', f"J{first}:J{last}")
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    return last


def build_not_shown_sheet(sheet: Worksheet) -> None:
    """Fill the sheet for what the customer does that was not shown."""
    sheet["A1"] = 'Ask last: "What do you do today that you did not see?"'
    sheet["A1"].font = Font(name=FONT, bold=True)
    write_header(sheet, NOT_SHOWN_FIRST - 1, NOT_SHOWN_TITLES)
    for row in range(NOT_SHOWN_FIRST, NOT_SHOWN_LAST + 1):
        for column in range(1, len(NOT_SHOWN_TITLES) + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font = Font(name=FONT)
            cell.alignment = WRAP
            cell.border = BORDER
            cell.fill = INPUT_FILL
        sheet.row_dimensions[row].height = 36
    for letter, width in zip("ABCDEF", NOT_SHOWN_WIDTHS, strict=True):
        sheet.column_dimensions[letter].width = width
    rows = f"{NOT_SHOWN_FIRST}:{{}}{NOT_SHOWN_LAST}"
    add_list(sheet, '"Yes,No,Partly"', "B" + rows.format("B"))
    add_list(sheet, OUTCOMES, "D" + rows.format("D"))
    add_list(sheet, '"Yes,No"', "F" + rows.format("F"))
    sheet.freeze_panes = f"A{NOT_SHOWN_FIRST}"


def build_how_to_sheet(sheet: Worksheet, last: int) -> None:
    """Fill the instructions, the example line and the progress counts."""
    first = HEADER_ROW + 1
    not_shown = f"'Not shown'!A{NOT_SHOWN_FIRST}:A{NOT_SHOWN_LAST}"
    lines: list[tuple[str | None, str | None]] = [
        ("Questions for the customer", None),
        (None, None),
        (
            "Which cells to fill",
            "The pale yellow cells. Everything else is the question list and "
            "should be left as it is.",
        ),
        (
            "One copy per customer",
            "Save a copy of this file with the customer's name before the meeting.",
        ),
        (
            "When to ask",
            "Ask section 1 before the demonstration: the answers decide what "
            "to show. Ask the rest after it.",
        ),
        ("Customer's answer", "Write the answer in the customer's own words."),
        (
            "Priority",
            "Critical = the answer can change how the application is designed, "
            "or decides whether it fits and how the firm is created; do not "
            "leave without it. High = may be something we have to build or "
            "support, or a setting needed before the first bill. Low = a "
            "setting that can wait. Filter this column when time is short.",
        ),
        (
            "Check first? = Yes",
            "The answer may be something the application does not do today. "
            'Do not say "yes" to it in the meeting.',
        ),
        (
            "Not shown sheet",
            'Ask last: "What do you do today that you did not see?" One line '
            "for each thing.",
        ),
        (
            "Outcome (at review)",
            "Setting = set up from the answer. Backlog item = to be built. "
            "Not done today = tell the customer plainly. To check = somebody "
            "must find out. No action = nothing to do.",
        ),
        (
            "Reviewed?",
            "Mark Yes when the line has been gone through. Filter this column "
            "on blanks to see what is left.",
        ),
        (None, None),
        (
            "Example of a filled line",
            "Ask: Do you give more discount for more quantity?  |  Customer's "
            'answer: "Yes, 2% extra above 10 cartons, 5% above 50."  |  '
            "Outcome: Setting  |  Review notes: Quantity breaks on the Retail "
            "price list  |  Reviewed?: Yes",
        ),
        (None, None),
        ("Progress", None),
        ("Questions", f"=COUNTA(Questions!C{first}:C{last})"),
        ("Answered", f"=COUNTA(Questions!G{first}:G{last})"),
        ("Critical questions", f'=COUNTIF(Questions!E{first}:E{last},"Critical")'),
        (
            "Critical answered",
            f'=COUNTIFS(Questions!E{first}:E{last},"Critical",'
            f'Questions!G{first}:G{last},"<>")',
        ),
        ("Reviewed", f'=COUNTIF(Questions!J{first}:J{last},"Yes")'),
        ("Marked Check first", f'=COUNTIF(Questions!F{first}:F{last},"Yes")'),
        ("Not-shown lines written", f"=COUNTA({not_shown})"),
        (None, None),
        (
            "Source",
            "docs/CUSTOMER_REQUIREMENTS_QUESTIONS.md. Change the questions "
            "there and rebuild this file with "
            "backend/scripts/make_customer_questions_xlsx.py.",
        ),
    ]
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 90
    for row, (label, text) in enumerate(lines, 1):
        left = sheet.cell(row=row, column=1, value=label)
        right = sheet.cell(row=row, column=2, value=text)
        left.font = Font(name=FONT, bold=True, size=14 if row == 1 else 10)
        right.font = Font(name=FONT)
        left.alignment = WRAP
        right.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")


def build_workbook(questions: list[Question]) -> Workbook:
    """Return the three-sheet workbook for the given questions."""
    workbook = Workbook()
    last = build_questions_sheet(workbook.active, questions)
    build_not_shown_sheet(workbook.create_sheet("Not shown"))
    build_how_to_sheet(workbook.create_sheet("How to use", 0), last)
    return workbook


def main() -> None:
    """Read the questions and write the workbook."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    questions = read_questions(args.source)
    if not questions:
        raise SystemExit(f"No questions found in {args.source}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(questions).save(args.output)
    print(f"{len(questions)} questions -> {args.output}")


if __name__ == "__main__":
    main()
