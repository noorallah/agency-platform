"""TC-MAST-027: a product import with no switch columns takes its category's goods type."""
import io
import pathlib, sys; sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from _gt import *  # noqa: F401,F403

c = Check("tc_mast_027")
tag = suffix()
admin = client("admin")
med = ensure_use(admin, "MEDICINE")
paint = ensure_use(admin, "PAINT")
tablets = cat(admin, tag, "TABLETS", goods_type_id=med["id"])
emulsions = cat(admin, tag, "EMULSIONS", goods_type_id=paint["id"])
sundries = cat(admin, tag, "SUNDRIES")
strips = cat(admin, tag, "STRIPS", parent_id=tablets["id"])


def by_code(code):
    status, body = admin.get(f"{PRODUCTS}?search={code}&page_size=20")
    rows = [r for r in data(body) if r["code"] == code]
    return rows[0] if rows else None


def code(name):
    return f"{tag}-{name}"


# The template: three optional tracking columns and a UnitSet column.
status, raw = admin.raw(f"{PRODUCTS}/import-template?format=xlsx")
c.eq(status, 200, "template downloads")
try:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(raw))
    notes = {row[0]: row for row in book["Notes"].iter_rows(values_only=True) if row and row[0]}
    for heading in ("TrackBatch", "TrackExpiry", "TrackSerial"):
        row = notes.get(heading)
        c.ok(row is not None and row[1] == "No" and "Blank takes the goods type" in str(row[2]),
             f"{heading} is optional and says blank takes the goods type", row)
    c.ok("UnitSet" in notes, "template has a UnitSet column")
except ImportError:
    c.ok(False, "openpyxl is not available to read the template")

cols1 = ["Code", "Name", "Category", "SubCategory"]
rows1 = [
    {"Code": code("IM-MED"), "Name": "Import med", "Category": tablets["code"]},
    {"Code": code("IM-SUB"), "Name": "Import sub", "Category": tablets["code"], "SubCategory": strips["code"]},
    {"Code": code("IM-PAINT"), "Name": "Import paint", "Category": emulsions["code"]},
    {"Code": code("IM-GEN"), "Name": "Import general", "Category": sundries["code"]},
]
status, body = import_file(admin, csv_bytes(rows1, cols1))
rep = data(body) if status == 200 else {}
c.eq((status, rep.get("issues"), rep.get("to_create")), (200, [], 4), f"first check is clean ({issues_text(body)})")
status, body = import_file(admin, csv_bytes(rows1, cols1), apply=True)
c.eq((status, data(body).get("imported") if status == 200 else None), (200, True), f"first file imports ({issues_text(body)})")


def sw(row):
    return {k: row.get(k) for k in ("track_batch", "track_expiry", "track_manufacturing_date", "track_serial")}


expect = {
    "IM-MED": dict(track_batch=True, track_expiry=True, track_manufacturing_date=True, track_serial=False),
    "IM-SUB": dict(track_batch=True, track_expiry=True, track_manufacturing_date=True, track_serial=False),
    "IM-PAINT": dict(track_batch=True, track_expiry=False, track_manufacturing_date=False, track_serial=False),
    "IM-GEN": dict(track_batch=False, track_expiry=False, track_manufacturing_date=False, track_serial=False),
}
for name, want in expect.items():
    row = by_code(code(name))
    c.eq(sw(row) if row else None, want, f"{name} switches")

cols2 = ["Code", "Name", "Category", "TrackBatch", "TrackExpiry", "TrackSerial"]
rows2 = [
    {"Code": code("IM-NO"), "Name": "Import no", "Category": tablets["code"], "TrackExpiry": "No"},
    {"Code": code("IM-YES"), "Name": "Import yes", "Category": sundries["code"], "TrackBatch": "Yes"},
    {"Code": code("IM-BAD"), "Name": "Import bad", "Category": sundries["code"], "TrackBatch": "maybe"},
]
status, body = import_file(admin, csv_bytes(rows2, cols2))
rep = data(body)
bad = [i for i in rep.get("issues", []) if i.get("code") == code("IM-BAD")]
c.ok(len(bad) == 1 and bad[0].get("column") == "TrackBatch", "IM-BAD is named by row and column", rep.get("issues"))
status, body = import_file(admin, csv_bytes(rows2, cols2), apply=True)
c.eq(data(body).get("imported"), False, "the file does not import while IM-BAD is in it")
c.eq(by_code(code("IM-NO")), None, "nothing was written by the refused apply")
status, body = import_file(admin, csv_bytes(rows2[:2], cols2), apply=True)
c.eq(data(body).get("imported"), True, f"the file imports without IM-BAD ({issues_text(body)})")
no = by_code(code("IM-NO"))
c.eq(sw(no) if no else None,
     dict(track_batch=True, track_expiry=False, track_manufacturing_date=True, track_serial=False),
     "IM-NO: a cell saying No wins; the rest follows Medicine")
yes = by_code(code("IM-YES"))
c.eq(sw(yes) if yes else None,
     dict(track_batch=True, track_expiry=False, track_manufacturing_date=False, track_serial=False),
     "IM-YES: the file's Yes, nothing else")

# A user without PRODUCT_IMPORT.
sales = client("sales")
c.eq(import_file(sales, csv_bytes(rows2[:1], cols2))[0], 403, "sales manager: import check is 403")
c.eq(sales.raw(f"{PRODUCTS}/import-template?format=csv")[0], 403, "sales manager: template is 403")
c.done()
