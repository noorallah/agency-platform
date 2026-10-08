"""TC-MAST-028: the UnitSet column fills a new product's units and its pack rule."""
import io
import pathlib, sys; sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from _gt import *  # noqa: F401,F403

c = Check("tc_mast_028")
tag = suffix()
admin = client("admin")
units = uom_ids(admin)
med = ensure_use(admin, "MEDICINE")
paint = ensure_use(admin, "PAINT")
tablets = cat(admin, tag, "TABLETS", goods_type_id=med["id"])
emulsions = cat(admin, tag, "EMULSIONS", goods_type_id=paint["id"])


def code(name):
    return f"{tag}-{name}"


def by_code(c_):
    rows = [r for r in data(admin.get(f"{PRODUCTS}?search={c_}&page_size=20")[1]) if r["code"] == c_]
    return rows[0] if rows else None


def rule_of(product_id):
    rows = data(admin.get(f"/api/v1/uom-framework/conversion-rules?product_id={product_id}&page_size=50")[1])
    return [(r["from_uom_id"], r["to_uom_id"], float(r["conversion_factor"])) for r in rows]


# The Lists sheet carries a Unit set column with every set offered.
status, raw = admin.raw(f"{PRODUCTS}/import-template?format=xlsx")
try:
    from openpyxl import load_workbook
    lists = load_workbook(io.BytesIO(raw))["Lists"]
    header = [cell.value for cell in lists[1]]
    index = header.index("Unit set")
    names = [row[index].value for row in lists.iter_rows(min_row=2) if row[index].value]
    wanted = [r["name"] for r in data(admin.get(SETS)[1]) if r["is_active"]]
    c.eq(sorted(names), sorted(wanted), "Lists sheet unit-set column holds every set offered")
except (ImportError, ValueError) as error:
    c.ok(False, f"could not read the Lists sheet: {error}")

# One existing product with its own units.
status, body = prod(admin, tag, "UX-OLD", category_id=tablets["id"], base_uom_id=units["PIECE"],
                    inventory_uom_id=units["PIECE"], purchase_uom_id=units["PIECE"])
c.eq(status, 201, f"UX-OLD created ({message(body)})")

cols = ["Code", "Name", "Category", "Unit", "UnitSet"]
r = lambda n, cat_, ucell="", uset="": {"Code": code(n), "Name": n, "Category": cat_["code"], "Unit": ucell, "UnitSet": uset}  # noqa: E731
file_rows = [
    r("UX-1", tablets, "", "Strip, box of 10"),
    r("UX-2", tablets, "BOX", "Strip, box of 10"),
    r("UX-3", tablets, "", "Blister, box of 99"),
    r("UX-4", emulsions, "", "Strip, box of 10"),
    # CASE TEXT: the shared catalogue has no "Tin, loose"; "Piece, loose" is the set tied to no goods type.
    r("UX-5", emulsions, "", "Piece, loose"),
    r("UX-OLD", tablets, "", "Strip, box of 10"),
]
status, body = import_file(admin, csv_bytes(file_rows, cols), existing="update")
rep = data(body)
c.eq([(i["code"], i["column"]) for i in rep.get("issues", [])], [(code("UX-3"), "UnitSet")], "only UX-3 is a problem, on UnitSet")
c.ok("is not an active unit set." in issues_text(body), "the message says it is not an active unit set", issues_text(body))
status, body = import_file(admin, csv_bytes(file_rows, cols), apply=True, existing="update")
c.eq(data(body).get("imported"), False, "nothing imported while row 3 is in the file")
c.eq(by_code(code("UX-1")), None, "UX-1 was not written by the refused apply")

good = [row for row in file_rows if row["Code"] != code("UX-3")]
status, body = import_file(admin, csv_bytes(good, cols), existing="update")
rep = data(body)
c.eq(rep.get("issues"), [], "after row 3 is removed the check has no problems")
warns = [(w["code"], w["message"]) for w in rep.get("warnings", [])]
c.eq(len(warns), 3, f"three warnings {warns}")
joined = " | ".join(m for _, m in warns)
c.ok("is marked for other goods types than this product's. It is imported as written." in joined, "UX-4 warning", joined)
c.ok("is passed over: a unit set fills a new product only, and this product keeps its units." in joined, "UX-OLD warning", joined)
c.ok("is passed over: the unit set 'Strip, box of 10' fills this product's units and its pack conversion." in joined, "UX-2 warning: its Unit is passed over", joined)
status, body = import_file(admin, csv_bytes(good, cols), apply=True, existing="update")
c.eq(data(body).get("imported"), True, f"imports ({issues_text(body)})")

ux1, ux2, ux4, ux5, old = (by_code(code(n)) for n in ("UX-1", "UX-2", "UX-4", "UX-5", "UX-OLD"))
if ux1:
    c.eq((ux1["base_uom_id"], ux1["purchase_uom_id"], ux1["inventory_uom_id"]), (units["STRIP"], units["BOX"], units["STRIP"]), "UX-1 units")
    c.eq(rule_of(ux1["id"]), [(units["BOX"], units["STRIP"], 10.0)], "UX-1 rule 1 Box = 10 Strip")
if ux2:
    c.eq((ux2["base_uom_id"], ux2["purchase_uom_id"]), (units["STRIP"], units["BOX"]), "UX-2 takes the set's units; its own Unit is passed over")
    c.eq(rule_of(ux2["id"]), [(units["BOX"], units["STRIP"], 10.0)], "UX-2 same conversion")
if ux4:
    c.eq(ux4["base_uom_id"], units["STRIP"], "UX-4 written with the Strip set")
if ux5:
    c.eq(rule_of(ux5["id"]), [], "UX-5 (Piece, loose) has no conversion")
if old:
    c.eq((old["base_uom_id"], rule_of(old["id"])), (units["PIECE"], []), "UX-OLD keeps its units and gains no rule")

# The Pack size heading is read as the UnitSet column.
cols2 = ["Code", "Name", "Category", "Pack size"]
status, body = import_file(admin, csv_bytes([{"Code": code("UX-6"), "Name": "UX-6", "Category": tablets["code"],
                                              "Pack size": "Strip, box of 15"}], cols2), apply=True)
rep = data(body)
c.eq(rep.get("imported"), True, f"Pack size file imports ({issues_text(body)})")
c.ok("UnitSet" in rep.get("columns_used", []), "Pack size is read as UnitSet", rep.get("columns_used"))
ux6 = by_code(code("UX-6"))
if ux6:
    c.eq(rule_of(ux6["id"]), [(units["BOX"], units["STRIP"], 15.0)], "UX-6 rule 15 from the Pack size heading")
c.done()
